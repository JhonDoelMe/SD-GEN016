import io
import csv
import datetime
from typing import Optional, List
from sqlalchemy import select, func, or_, and_
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.app.core.timezone import to_facility_time, now_facility
from backend.app.models.generator import Generator, GeneratorRun
from backend.app.models.fuel import FuelStock, FuelReceipt, FuelTransfer
from backend.app.models.maintenance import MaintenanceRecord, MaintenanceSchedule
from backend.app.models.fault import Fault
from backend.app.models.audit import AuditLog, SystemAdjustment
from backend.app.models.facility import Facility
from backend.app.schemas.report import (
    OperationalReportOut,
    ReportRunItem,
    ReportFuelReceiptItem,
    ReportFuelTransferItem,
    ReportMaintenanceItem,
)


def utc_now() -> datetime.datetime:
    return datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None)


def parse_report_dt(dt_str: Optional[str], is_end: bool = False) -> Optional[datetime.datetime]:
    if not dt_str:
        return None
    try:
        dt = datetime.datetime.fromisoformat(dt_str)
    except ValueError:
        return None
    if dt.tzinfo is not None:
        dt = dt.astimezone(datetime.timezone.utc).replace(tzinfo=None)
    if is_end and len(dt_str) <= 10:
        dt = dt.replace(hour=23, minute=59, second=59, microsecond=999999)
    return dt


def fmt_dt(dt: Optional[datetime.datetime]) -> str:
    if not dt:
        return "-"
    return to_facility_time(dt).strftime("%d.%m.%Y, %H:%M:%S")


def fmt_date(dt: Optional[datetime.datetime]) -> str:
    if not dt:
        return "-"
    return to_facility_time(dt).strftime("%d.%m.%Y")


def fmt_time(dt: Optional[datetime.datetime]) -> str:
    if not dt:
        return "-"
    return to_facility_time(dt).strftime("%H:%M:%S")


def fmt_duration(seconds: Optional[int], hours: Optional[float]) -> str:
    if seconds is not None and seconds >= 0:
        h = seconds // 3600
        m = (seconds % 3600) // 60
        s = seconds % 60
        return f"{h:02d}:{m:02d}:{s:02d}"
    if hours is not None and hours >= 0:
        tot_sec = int(round(hours * 3600))
        h = tot_sec // 3600
        m = (tot_sec % 3600) // 60
        s = tot_sec % 60
        return f"{h:02d}:{m:02d}:{s:02d}"
    return "00:00:00"


def fmt_duration_human(seconds: Optional[int], hours: Optional[float]) -> str:
    if seconds is None and hours is not None:
        seconds = int(round(hours * 3600))
    if seconds is None or seconds <= 0:
        return "0 хв"
    h = seconds // 3600
    m = (seconds % 3600) // 60
    s = seconds % 60
    if h > 0:
        if m > 0:
            return f"{h} год {m:02d} хв"
        return f"{h} год"
    if m > 0:
        if s > 0:
            return f"{m} хв {s:02d} с"
        return f"{m} хв"
    return f"{s} с"


MONTHS_UA_SHORT = {
    1: "січ.", 2: "лют.", 3: "бер.", 4: "квіт.", 5: "трав.", 6: "черв.",
    7: "лип.", 8: "серп.", 9: "вер.", 10: "жовт.", 11: "лист.", 12: "груд."
}


def fmt_dt_ua(dt: Optional[datetime.datetime]) -> str:
    if not dt:
        return "-"
    d = to_facility_time(dt)
    m = MONTHS_UA_SHORT.get(d.month, "")
    return f"{d.day} {m} {d.year} р., {d.strftime('%H:%M:%S')}"


def fmt_date_ua(dt: Optional[datetime.datetime]) -> str:
    if not dt:
        return "-"
    d = to_facility_time(dt)
    m = MONTHS_UA_SHORT.get(d.month, "")
    return f"{d.day} {m} {d.year} р."


def fmt_hours_human(hours: Optional[float]) -> str:
    if hours is None or hours <= 0:
        return "0 хв"
    total_minutes = int(round(hours * 60))
    h = total_minutes // 60
    m = total_minutes % 60
    if h == 0:
        return f"{m} хв"
    return f"{h} год {m:02d} хв"



def sanitize_spreadsheet_cell(val):
    """Sanitize string values against CSV/Excel Formula Injection (=, +, -, @, tab, CR)."""
    if isinstance(val, str) and val and val[0] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{val}"
    return val


def auto_adjust_column_widths(ws, min_width=12, max_width=50, padding=3):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if "\n" in val_str:
                lines = val_str.split("\n")
                max_len = max(max_len, max(len(l) for l in lines))
            else:
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + padding, min_width), max_width)


# Styling Definitions
PRIMARY_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")     # Deep Navy
SECONDARY_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")   # Blue
SECTION_FILL = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")     # Teal
ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")       # Light Slate
TOTAL_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")       # Total grey

HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
SECTION_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
BOLD_FONT = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
REGULAR_FONT = Font(name="Segoe UI", size=10, color="0F172A")
TOTAL_FONT = Font(name="Segoe UI", size=10, bold=True, color="1E3A8A")

THIN_BORDER_SIDE = Side(border_style="thin", color="CBD5E1")
DOUBLE_BOTTOM_SIDE = Side(border_style="double", color="1E3A8A")

CELL_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)
TOTAL_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=DOUBLE_BOTTOM_SIDE)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

# Українські мапування статусів та системних дій для звітів Excel
GENERATOR_STATUS_UA = {
    "STOPPED": "Зупинено",
    "RUNNING": "В роботі",
    "MAINTENANCE": "На обслуговуванні",
    "FAULT": "Несправність / Аварія",
}

RUN_STATUS_UA = {
    "ACTIVE": "В роботі",
    "COMPLETED": "Завершено",
}

FAULT_STATUS_UA = {
    "OPEN": "Відкрито",
    "IN_PROGRESS": "В роботі",
    "RESOLVED": "Усунено",
    "CLOSED": "Закрито",
}

FAULT_PRIORITY_UA = {
    "LOW": "Низький",
    "MEDIUM": "Середній",
    "HIGH": "Високий",
    "CRITICAL": "Критичний",
}

AUDIT_ACTION_UA = {
    # Generator
    "GENERATOR_START": "Запуск генератора",
    "GENERATOR_STOP": "Зупинка генератора",
    "GENERATOR_CREATED": "Створення генератора",
    "GENERATOR_UPDATED": "Оновлення параметрів генератора",
    "GENERATOR_DELETED": "Видалення генератора",
    "GENERATOR_MANUAL_RUN_ADDED": "Ручне внесення мотогодин",
    "GENERATOR_WIZARD_CONFIGURED": "Первинне налаштування генератора",

    # Fuel
    "FUEL_RECEIPT": "Оприбуткування палива на склад",
    "FUEL_RECEIPT_ADDED": "Оприбуткування палива на склад",
    "FUEL_TRANSFER_TO_TANK": "Заправка бака зі складу",
    "FUEL_TRANSFERRED": "Заправка бака зі складу",
    "FUEL_STOCK_CREATED": "Створення складу ГСМ",
    "FUEL_STOCK_UPDATED": "Оновлення складу ГСМ",
    "FUEL_STOCK_DELETED": "Видалення складу ГСМ",
    "FUEL_PRICES_REFRESHED": "Оновлення цін АЗК WOG",

    # Maintenance
    "MAINTENANCE_RECORDED": "Запис ТО",
    "MAINTENANCE_INTERMEDIATE": "Проміжне ТО",
    "MAINTENANCE_SCHEDULED": "Регламентне ТО",
    "MAINTENANCE_SCHEDULE_RECALCULATED": "Перерахунок графіка ТО",
    "SCHEDULE_RECALCULATED": "Перерахунок графіка ТО",

    # Faults
    "FAULT_REPORTED": "Повідомлення про несправність",
    "FAULT_STATUS_UPDATED": "Зміна статусу несправності",
    "FAULT_RESOLVED": "Усунення несправності",

    # Auth
    "AUTH_LOGIN_SUCCESS": "Успішний вхід у систему",
    "AUTH_LOGIN_FAILED": "Невдала спроба входу",
    "AUTH_LOGIN_LOCKED": "Тимчасове блокування входу",
    "AUTH_LOGIN_BLOCKED": "Блокування входу",
    "AUTH_LOGOUT": "Вихід із системи",
    "USER_LOGIN": "Вхід у систему",
    "USER_LOGOUT": "Вихід із системи",

    # Users
    "USER_CREATED": "Створення користувача",
    "USER_UPDATED": "Редагування користувача",
    "USER_DELETED": "Видалення користувача",
    "USER_PASSWORD_RESET": "Скидання пароля користувача",

    # System & Adjustments
    "SYSTEM_ADJUSTMENT": "Системне коригування лічильників",
    "ADJUSTMENT_APPLIED": "Системне коригування лічильників",
    "NOTIFICATION_SETTINGS_UPDATED": "Оновлення налаштувань сповіщень",
    "SYSTEM_SETTINGS_UPDATED": "Оновлення системних налаштувань",
    "SYSTEM_SETTINGS_RESET": "Скидання системних налаштувань",
    "SYSTEM_FACILITY_CLEAN": "Очищення даних об'єкта",
    "SYSTEM_REPORTS_CLEAN": "Очищення звітів",
    "SYSTEM_FACTORY_RESET": "Повне заводське скидання",
    "SYSTEM_RESET": "Очищення / скидання даних",

    # Facility
    "FACILITY_CREATED": "Створення об'єкта",
    "FACILITY_UPDATED": "Оновлення об'єкта",
    "FACILITY_DELETED": "Видалення об'єкта",
}

AUDIT_ENTITY_UA = {
    "Generator": "Генератор",
    "Facility": "Об'єкт",
    "FuelStock": "Склад ГСМ",
    "FuelReceipt": "Оприбуткування",
    "FuelTransfer": "Заправка",
    "MaintenanceRecord": "Запис ТО",
    "MaintenanceSchedule": "Графік ТО",
    "Fault": "Несправність",
    "User": "Користувач",
    "SystemSetting": "Системні налаштування",
}


def format_audit_details_ua(details_obj: Any, action: str = "") -> str:
    """Перетворює JSON деталей аудиту на зрозумілий людині україномовний опис для звіту."""
    import json
    if not details_obj:
        return ""
    if isinstance(details_obj, str):
        try:
            details = json.loads(details_obj)
        except Exception:
            return details_obj
    elif isinstance(details_obj, dict):
        details = details_obj
    else:
        return str(details_obj)

    if not isinstance(details, dict):
        return str(details)

    # 1. Зупинка генератора (підсумок роботи)
    if "duration_seconds" in details or "calculated_consumption_l" in details:
        parts = []
        if "duration_formatted" in details:
            parts.append(f"Час роботи: {details['duration_formatted']}")
        elif "duration_hours" in details:
            parts.append(f"Час роботи: {details['duration_hours']} год")
        if "calculated_consumption_l" in details:
            parts.append(f"Витрата: {details['calculated_consumption_l']} л")
        if "end_fuel_level_l" in details:
            parts.append(f"Залишок у баку: {details['end_fuel_level_l']} л")
        if "end_hours" in details:
            parts.append(f"Мотогодини: {details['end_hours']} год")
        return " | ".join(parts) if parts else str(details)

    # 2. Запуск генератора
    if "start_hours" in details and "start_fuel" in details:
        return f"Початкові мотогодини: {details.get('start_hours', '-')} год | Паливо в баку: {details.get('start_fuel', '-')} л"

    # 3. Системне коригування (Audit adjustments)
    if "field" in details:
        field_map = {
            "current_operating_hours": "Мотогодини",
            "fuel_tank_level_l": "Бак генератора",
            "stock_balance_l": "Залишок на складі",
        }
        field_name = field_map.get(details.get("field"), details.get("field"))
        old_val = details.get("old_value", "-")
        new_val = details.get("new_value", "-")
        reason = details.get("reason", "")
        res = f"Коригування '{field_name}': {old_val} ➔ {new_val}"
        if reason:
            res += f" (Причина: {reason})"
        return res

    # 4. Оприбуткування палива на склад
    if "receipt_id" in details and "cost_total" in details:
        st_name = details.get("stock_name", "")
        st_str = f" ({st_name})" if st_name else ""
        return f"Оприбутковано {details.get('liters', '-')} л на суму {details.get('cost_total', '-')} грн{st_str}"

    # 5. Заправка в бак зі складу
    if "liters" in details and ("stock_before" in details or "tank_before" in details):
        return f"Заправлено в бак: {details.get('liters')} л"

    # 6. Керування користувачами
    if "login" in details and ("roles" in details or action in ("USER_CREATED", "USER_UPDATED", "USER_DELETED")):
        role_map = {
            "superadmin": "Суперадміністратор",
            "admin": "Адміністратор",
            "operator": "Оператор",
            "viewer": "Спостерігач",
        }
        login = details.get("login")
        roles = details.get("roles", [])
        if roles:
            roles_ua = ", ".join(role_map.get(r, r) for r in roles)
            return f"Користувач: {login} (Роль: {roles_ua})"
        return f"Користувач: {login}"

    # 7. Налаштування сповіщень
    if "updated_keys" in details:
        return "Оновлено параметри та перемикачі сповіщень"

    # 8. Перерахунок регламенту ТО
    if "next_due_hours" in details and "interval_hours" in details:
        return f"Мотогодини: {details.get('current_hours', '-')} год | Наступне планове ТО: {details.get('next_due_hours')} год"

    # 9. Запис ТО
    if "work" in details and "type" in details:
        m_type = "Регламентне ТО" if details.get("type") == "SCHEDULED" else "Проміжне ТО"
        return f"{m_type} | Роботи: {details.get('work', '')} | Вартість: {details.get('cost', 0)} грн"

    # 10. Загальний переклад словника
    key_translations = {
        "title": "Назва",
        "description": "Опис",
        "priority": "Пріоритет",
        "status": "Статус",
        "operating_hours": "Мотогодини",
        "generator_status_after": "Статус генератора",
        "liters": "Літри",
        "cost": "Вартість",
        "cost_total": "Загальна сума",
        "price_per_l": "Ціна за літр",
        "source": "Джерело",
        "updated_at": "Оновлено",
        "name": "Назва",
        "tank_capacity_l": "Місткість бака",
        "initial_hours": "Початкові мотогодини",
        "initial_balance_l": "Початковий залишок",
        "facility_timezone": "Часовий пояс",
        "work_schedule_start": "Початок зміни",
        "work_schedule_end": "Кінець зміни",
        "organization_name": "Організація",
        "work": "Виконані роботи",
        "type": "Тип",
        "fuel_type": "Тип палива",
        "field": "Поле",
        "old_value": "Старе значення",
        "new_value": "Нове значення",
        "reason": "Причина",
    }
    val_translations = {
        "STOPPED": "Зупинено",
        "RUNNING": "В роботі",
        "MAINTENANCE": "На обслуговуванні",
        "FAULT": "Несправність",
        "LOW": "Низький",
        "MEDIUM": "Середній",
        "HIGH": "Високий",
        "CRITICAL": "Критичний",
        "OPEN": "Відкрито",
        "RESOLVED": "Усунено",
        "CLOSED": "Закрито",
        "SCHEDULED": "Регламентне ТО",
        "INTERMEDIATE": "Проміжне ТО",
    }

    parts = []
    for k, v in details.items():
        if k in ("facility_id", "generator_id", "run_id", "stock_id", "user_id", "receipt_id"):
            continue
        k_ua = key_translations.get(k, k)
        if isinstance(v, str):
            v_ua = val_translations.get(v.upper(), v)
            if len(v_ua) == 19 and v_ua[4] == '-' and v_ua[7] == '-' and v_ua[10] == ' ':
                try:
                    dt_p = datetime.datetime.strptime(v_ua, "%Y-%m-%d %H:%M:%S")
                    v_ua = dt_p.strftime("%d.%m.%Y, %H:%M:%S")
                except Exception:
                    pass
        elif isinstance(v, dict):
            sub_parts = [f"{key_translations.get(sk, sk)}: {val_translations.get(str(sv).upper(), sv)}" for sk, sv in v.items()]
            v_ua = "{" + ", ".join(sub_parts) + "}"
        else:
            v_ua = str(v)
        parts.append(f"{k_ua}: {v_ua}")
    return " | ".join(parts) if parts else str(details)


async def generate_operational_report(
    db: AsyncSession,
    generator_id: Optional[int] = None,
    facility_id: Optional[int] = None,
    start_date: Optional[datetime.datetime] = None,
    end_date: Optional[datetime.datetime] = None
) -> OperationalReportOut:
    gen_query = select(Generator).options(selectinload(Generator.facility))
    if generator_id:
        gen_query = gen_query.where(Generator.id == generator_id)
    elif facility_id:
        gen_query = gen_query.where(Generator.facility_id == facility_id)
    gen_res = await db.execute(gen_query)
    generator = gen_res.scalars().first()
    gen_name = generator.name if generator else "Генератор"
    facility_name = generator.facility.name if generator and generator.facility else None
    if not facility_name and facility_id:
        fac = (await db.execute(select(Facility).where(Facility.id == facility_id))).scalars().first()
        if fac:
            facility_name = fac.name

    now = utc_now()
    if not end_date:
        end_date = now
    if not start_date:
        start_date = end_date - datetime.timedelta(days=30)

    # 1. Runs aggregate
    runs_q = select(
        func.count(GeneratorRun.id),
        func.coalesce(func.sum(GeneratorRun.duration_hours), 0.0),
        func.coalesce(func.sum(GeneratorRun.calculated_consumption_l), 0.0)
    ).where(
        GeneratorRun.start_time >= start_date,
        GeneratorRun.start_time <= end_date,
        GeneratorRun.status == "COMPLETED"
    )
    if generator:
        runs_q = runs_q.where(GeneratorRun.generator_id == generator.id)
    runs_res = await db.execute(runs_q)
    total_runs, total_hours, total_cons = runs_res.first()

    # 2. Fuel receipts aggregate
    rec_q = select(
        func.coalesce(func.sum(FuelReceipt.liters), 0.0),
        func.coalesce(func.sum(FuelReceipt.cost_total), 0.0)
    ).where(
        FuelReceipt.created_at >= start_date,
        FuelReceipt.created_at <= end_date
    )
    if facility_id:
        rec_q = rec_q.outerjoin(FuelStock, FuelReceipt.stock_id == FuelStock.id).where(
            or_(
                FuelReceipt.facility_id == facility_id,
                FuelStock.facility_id == facility_id,
                and_(facility_id == 1, FuelReceipt.facility_id.is_(None), FuelStock.facility_id.is_(None))
            )
        )
    rec_res = await db.execute(rec_q)
    fuel_rec_l, fuel_rec_cost = rec_res.first()
    avg_price = round(fuel_rec_cost / fuel_rec_l, 2) if fuel_rec_l > 0 else 0.0

    # 3. Fuel transfers aggregate
    tr_q = select(
        func.coalesce(func.sum(FuelTransfer.liters), 0.0)
    ).where(
        FuelTransfer.created_at >= start_date,
        FuelTransfer.created_at <= end_date
    )
    if generator:
        tr_q = tr_q.where(FuelTransfer.generator_id == generator.id)
    elif facility_id:
        tr_q = tr_q.where(FuelTransfer.facility_id == facility_id)
    tr_res = await db.execute(tr_q)
    fuel_transferred_l = tr_res.scalar_one()

    # Stock & tank
    if facility_id:
        stock_bal_q = select(func.coalesce(func.sum(FuelStock.current_balance_l), 0.0)).where(FuelStock.facility_id == facility_id)
        stock_bal = (await db.execute(stock_bal_q)).scalar_one()
    else:
        stock_bal_q = select(func.coalesce(func.sum(FuelStock.current_balance_l), 0.0))
        stock_bal = (await db.execute(stock_bal_q)).scalar_one()
    tank_bal = generator.fuel_tank_level_l if generator else 0.0

    # 4. Maintenance aggregate
    m_all_q = select(func.count(MaintenanceRecord.id), func.coalesce(func.sum(MaintenanceRecord.cost), 0.0)).where(
        MaintenanceRecord.created_at >= start_date,
        MaintenanceRecord.created_at <= end_date
    )
    if generator:
        m_all_q = m_all_q.where(MaintenanceRecord.generator_id == generator.id)
    m_all_res = await db.execute(m_all_q)
    m_count, m_cost = m_all_res.first()

    m_sched_q = select(func.count(MaintenanceRecord.id)).where(
        MaintenanceRecord.created_at >= start_date,
        MaintenanceRecord.created_at <= end_date,
        MaintenanceRecord.maintenance_type == "SCHEDULED"
    )
    if generator:
        m_sched_q = m_sched_q.where(MaintenanceRecord.generator_id == generator.id)
    m_sched_cnt = (await db.execute(m_sched_q)).scalar_one()
    m_inter_cnt = m_count - m_sched_cnt

    # 5. Faults aggregate
    f_total_q = select(func.count(Fault.id)).where(
        Fault.created_at >= start_date,
        Fault.created_at <= end_date
    )
    if generator:
        f_total_q = f_total_q.where(Fault.generator_id == generator.id)
    f_total = (await db.execute(f_total_q)).scalar_one()

    f_res_q = select(func.count(Fault.id)).where(
        Fault.created_at >= start_date,
        Fault.created_at <= end_date,
        Fault.status.in_(["RESOLVED", "CLOSED"])
    )
    if generator:
        f_res_q = f_res_q.where(Fault.generator_id == generator.id)
    f_resolved = (await db.execute(f_res_q)).scalar_one()

    # 6. Detailed records for the period
    # 6a. Generator runs
    runs_detail_q = select(GeneratorRun).options(selectinload(GeneratorRun.user)).where(
        GeneratorRun.start_time >= start_date,
        GeneratorRun.start_time <= end_date,
        GeneratorRun.status == "COMPLETED"
    )
    if generator:
        runs_detail_q = runs_detail_q.where(GeneratorRun.generator_id == generator.id)
    elif facility_id:
        runs_detail_q = runs_detail_q.join(Generator).where(Generator.facility_id == facility_id)
    runs_detail_q = runs_detail_q.order_by(GeneratorRun.start_time.desc()).limit(200)
    runs_rows = (await db.execute(runs_detail_q)).scalars().all()
    runs_list = [
        ReportRunItem(
            id=r.id,
            start_time=fmt_dt_ua(r.start_time),
            end_time=fmt_dt_ua(r.end_time) if r.end_time else None,
            duration_hours=round(r.duration_hours or 0.0, 2),
            duration_formatted=fmt_duration(r.duration_seconds, r.duration_hours),
            duration_human=fmt_hours_human(r.duration_hours),
            calculated_consumption_l=round(r.calculated_consumption_l or 0.0, 2),
            start_hours=round(r.start_hours or 0.0, 2),
            end_hours=round(r.end_hours or 0.0, 2) if r.end_hours is not None else None,
            operator_name=(r.user.full_name or r.user.login) if r.user else "Система",
            status=r.status or "COMPLETED"
        )
        for r in runs_rows
    ]

    # 6b. Fuel receipts
    rec_detail_q = select(FuelReceipt).options(
        selectinload(FuelReceipt.user),
        selectinload(FuelReceipt.stock)
    ).where(
        FuelReceipt.created_at >= start_date,
        FuelReceipt.created_at <= end_date
    )
    if facility_id:
        rec_detail_q = rec_detail_q.outerjoin(FuelStock, FuelReceipt.stock_id == FuelStock.id).where(
            or_(
                FuelReceipt.facility_id == facility_id,
                FuelStock.facility_id == facility_id,
                and_(facility_id == 1, FuelReceipt.facility_id.is_(None), FuelStock.facility_id.is_(None))
            )
        )
    rec_detail_q = rec_detail_q.order_by(FuelReceipt.created_at.desc()).limit(200)
    receipts_rows = (await db.execute(rec_detail_q)).scalars().all()
    receipts_list = [
        ReportFuelReceiptItem(
            id=rec.id,
            created_at=fmt_dt_ua(rec.created_at),
            liters=round(rec.liters or 0.0, 2),
            price_per_liter=round(rec.price_per_liter or 0.0, 2),
            cost_total=round(rec.cost_total or 0.0, 2),
            receipt_number=rec.receipt_number or "-",
            driver_name=rec.driver_name or "-",
            supplier=rec.receipt_number or "-",
            stock_name=rec.stock.name if rec.stock else "Склад ГСМ",
            operator_name=(rec.user.full_name or rec.user.login) if rec.user else "Система"
        )
        for rec in receipts_rows
    ]

    # 6c. Fuel transfers
    tr_detail_q = select(FuelTransfer).options(
        selectinload(FuelTransfer.user),
        selectinload(FuelTransfer.stock),
        selectinload(FuelTransfer.generator)
    ).where(
        FuelTransfer.created_at >= start_date,
        FuelTransfer.created_at <= end_date
    )
    if generator:
        tr_detail_q = tr_detail_q.where(FuelTransfer.generator_id == generator.id)
    elif facility_id:
        tr_detail_q = tr_detail_q.outerjoin(Generator, FuelTransfer.generator_id == Generator.id).where(
            or_(FuelTransfer.facility_id == facility_id, Generator.facility_id == facility_id)
        )
    tr_detail_q = tr_detail_q.order_by(FuelTransfer.created_at.desc()).limit(200)
    transfers_rows = (await db.execute(tr_detail_q)).scalars().all()
    transfers_list = [
        ReportFuelTransferItem(
            id=tr.id,
            created_at=fmt_dt_ua(tr.created_at),
            liters=round(tr.liters or 0.0, 2),
            stock_name=tr.stock.name if tr.stock else "Склад ГСМ",
            generator_name=tr.generator.name if tr.generator else (gen_name or "Генератор"),
            operator_name=(tr.user.full_name or tr.user.login) if tr.user else "Система"
        )
        for tr in transfers_rows
    ]

    # 6d. Maintenance records
    m_detail_q = select(MaintenanceRecord).options(
        selectinload(MaintenanceRecord.user),
        selectinload(MaintenanceRecord.generator)
    ).where(
        MaintenanceRecord.created_at >= start_date,
        MaintenanceRecord.created_at <= end_date
    )
    if generator:
        m_detail_q = m_detail_q.where(MaintenanceRecord.generator_id == generator.id)
    m_detail_q = m_detail_q.order_by(MaintenanceRecord.created_at.desc()).limit(200)
    m_rows = (await db.execute(m_detail_q)).scalars().all()
    maintenance_list = [
        ReportMaintenanceItem(
            id=m.id,
            created_at=fmt_dt_ua(m.created_at),
            maintenance_type="Регламентне ТО" if m.maintenance_type == "SCHEDULED" else "Проміжне ТО",
            operating_hours=round(m.operating_hours or 0.0, 2),
            work_description=m.work_description or "-",
            cost=round(m.cost or 0.0, 2),
            technician=m.user.full_name if (m.user and m.user.full_name) else (m.comment or "-"),
            operator_name=(m.user.full_name or m.user.login) if m.user else "Система"
        )
        for m in m_rows
    ]

    # 7. Additional calculations
    total_costs = round((fuel_rec_cost or 0.0) + (m_cost or 0.0), 2)
    avg_cons = round(total_cons / total_hours, 2) if total_hours > 0 else 0.0
    avg_run_dur = round(total_hours / total_runs, 2) if total_runs > 0 else 0.0

    return OperationalReportOut(
        start_date=start_date.strftime("%Y-%m-%d"),
        end_date=end_date.strftime("%Y-%m-%d"),
        start_date_human=fmt_date_ua(start_date),
        end_date_human=fmt_date_ua(end_date),
        generator_name=gen_name,
        facility_name=facility_name,
        total_runs_count=total_runs,
        total_operating_hours=round(total_hours, 2),
        total_operating_hours_human=fmt_hours_human(total_hours),
        total_calculated_consumption_l=round(total_cons, 2),
        total_fuel_received_l=round(fuel_rec_l, 2),
        total_fuel_receipts_cost_uah=round(fuel_rec_cost, 2),
        avg_fuel_price_per_liter=avg_price,
        total_fuel_transferred_to_tank_l=round(fuel_transferred_l, 2),
        current_stock_balance_l=round(stock_bal, 2),
        current_tank_level_l=round(tank_bal, 2),
        total_maintenance_count=m_count,
        scheduled_maintenance_count=m_sched_cnt,
        intermediate_maintenance_count=m_inter_cnt,
        total_maintenance_cost_uah=round(m_cost, 2),
        total_faults_count=f_total,
        resolved_faults_count=f_resolved,
        open_faults_count=(f_total - f_resolved),
        total_costs_uah=total_costs,
        avg_consumption_l_per_hour=avg_cons,
        avg_run_duration_hours=avg_run_dur,
        avg_run_duration_human=fmt_hours_human(avg_run_dur),
        runs=runs_list,
        fuel_receipts=receipts_list,
        fuel_transfers=transfers_list,
        maintenance_records=maintenance_list
    )


def export_report_to_csv(report: OperationalReportOut) -> str:
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')

    writer.writerow(["Параметр", "Значення", "Одиниця виміру"])
    writer.writerow(["Період звіт", f"{report.start_date} - {report.end_date}", ""])
    writer.writerow(["Генератор", sanitize_spreadsheet_cell(report.generator_name), ""])
    writer.writerow(["Кількість робочих циклів (пусків)", report.total_runs_count, "шт"])
    writer.writerow(["Відпрацьовано мотогодин", report.total_operating_hours, "год"])
    writer.writerow(["Розрахункова витрата палива", report.total_calculated_consumption_l, "л"])
    writer.writerow(["Оприбутковано палива на склад", report.total_fuel_received_l, "л"])
    writer.writerow(["Загальна вартість закупленого палива", report.total_fuel_receipts_cost_uah, "грн"])
    writer.writerow(["Середня вартість літра", report.avg_fuel_price_per_liter, "грн/л"])
    writer.writerow(["Заправлено в бак генератора", report.total_fuel_transferred_to_tank_l, "л"])
    writer.writerow(["Поточний залишок на складі ГСМ", report.current_stock_balance_l, "л"])
    writer.writerow(["Поточний рівень у баку", report.current_tank_level_l, "л"])
    writer.writerow(["Кількість проведених ТО", report.total_maintenance_count, "шт"])
    writer.writerow(["- з них регламентних ТО", report.scheduled_maintenance_count, "шт"])
    writer.writerow(["- з них проміжних ТО", report.intermediate_maintenance_count, "шт"])
    writer.writerow(["Витрати на ТО та матеріали", report.total_maintenance_cost_uah, "грн"])
    writer.writerow(["Зареєстровано несправностей", report.total_faults_count, "шт"])
    writer.writerow(["Усунено несправностей", report.resolved_faults_count, "шт"])
    writer.writerow(["Відкритих несправностей", report.open_faults_count, "шт"])

    return output.getvalue()


async def generate_excel_report(
    db: AsyncSession,
    generator_id: Optional[int] = None,
    facility_id: Optional[int] = None,
    start_date: Optional[datetime.datetime] = None,
    end_date: Optional[datetime.datetime] = None
) -> bytes:
    """
    Генерує повнофункціональний, професійно оформлений звіт у форматі Excel (.xlsx).
    Містить 5 тематичних вкладок:
    1. Загальний звіт (Дашборд KPI, параметри об'єкта та генератора)
    2. Звіт по днях (Детальні цикли роботи, години, паливо, тривалість)
    3. Рух палива та Склад ГСМ (Надходження, заправки в бак, залишки)
    4. Технічне обслуговування (Журнал регламентних та проміжних ТО, витрати)
    5. Журнал несправностей та інцидентів (Облік несправностей, критичність та усунення)
    """
    now = utc_now()
    if not end_date:
        end_date = now
    if not start_date:
        start_date = end_date - datetime.timedelta(days=30)

    # 1. Fetch Facility & Generator
    gen_query = select(Generator).options(selectinload(Generator.facility))
    if generator_id:
        gen_query = gen_query.where(Generator.id == generator_id)
    elif facility_id:
        gen_query = gen_query.where(Generator.facility_id == facility_id)
    gen_res = await db.execute(gen_query)
    generator = gen_res.scalars().first()

    facility = None
    if generator and generator.facility:
        facility = generator.facility
    elif facility_id:
        fac_res = await db.execute(select(Facility).where(Facility.id == facility_id))
        facility = fac_res.scalars().first()
    else:
        fac_res = await db.execute(select(Facility).order_by(Facility.id.asc()))
        facility = fac_res.scalars().first()

    facility_name = facility.name if facility else "Основний об'єкт"
    generator_name = generator.name if generator else "Бензиновий генератор"
    gen_model = generator.model if generator else "-"
    gen_serial = generator.serial_number if generator else "-"

    # 2. Fetch Operational Report Metrics
    op_report = await generate_operational_report(
        db=db,
        generator_id=generator.id if generator else None,
        facility_id=facility.id if facility else None,
        start_date=start_date,
        end_date=end_date
    )

    # 3. Fetch Maintenance Schedule
    m_sched_q = select(MaintenanceSchedule)
    if generator:
        m_sched_q = m_sched_q.where(MaintenanceSchedule.generator_id == generator.id)
    m_sched = (await db.execute(m_sched_q)).scalars().first()
    interval_hours = m_sched.interval_hours if m_sched else 300.0
    last_maint_hours = m_sched.last_performed_hours if m_sched else 0.0
    next_due_hours = m_sched.next_due_hours if m_sched else 300.0
    current_hours = generator.current_operating_hours if generator else 0.0
    hours_remaining = round(next_due_hours - current_hours, 1)

    # 4. Fetch Generator Runs
    runs_q = select(GeneratorRun).options(selectinload(GeneratorRun.user)).where(
        GeneratorRun.start_time >= start_date,
        GeneratorRun.start_time <= end_date
    )
    if generator:
        runs_q = runs_q.where(GeneratorRun.generator_id == generator.id)
    elif facility:
        runs_q = runs_q.join(Generator).where(Generator.facility_id == facility.id)
    runs_q = runs_q.order_by(GeneratorRun.start_time.asc())
    runs = (await db.execute(runs_q)).scalars().all()

    # 5. Fetch Fuel Receipts
    rec_q = select(FuelReceipt).options(
        selectinload(FuelReceipt.user),
        selectinload(FuelReceipt.stock).selectinload(FuelStock.facility)
    ).where(
        FuelReceipt.created_at >= start_date,
        FuelReceipt.created_at <= end_date
    )
    if facility:
        rec_q = rec_q.outerjoin(FuelStock, FuelReceipt.stock_id == FuelStock.id).where(
            or_(
                FuelReceipt.facility_id == facility.id,
                FuelStock.facility_id == facility.id,
                and_(facility.id == 1, FuelReceipt.facility_id.is_(None), FuelStock.facility_id.is_(None))
            )
        )
    rec_q = rec_q.order_by(FuelReceipt.created_at.asc())
    receipts = (await db.execute(rec_q)).scalars().all()

    # 6. Fetch Fuel Transfers
    tr_q = select(FuelTransfer).options(
        selectinload(FuelTransfer.user),
        selectinload(FuelTransfer.stock).selectinload(FuelStock.facility),
        selectinload(FuelTransfer.generator)
    ).where(
        FuelTransfer.created_at >= start_date,
        FuelTransfer.created_at <= end_date
    )
    if generator:
        tr_q = tr_q.where(FuelTransfer.generator_id == generator.id)
    elif facility:
        tr_q = tr_q.outerjoin(Generator, FuelTransfer.generator_id == Generator.id).where(
            or_(
                FuelTransfer.facility_id == facility.id,
                Generator.facility_id == facility.id
            )
        )
    tr_q = tr_q.order_by(FuelTransfer.created_at.asc())
    transfers = (await db.execute(tr_q)).scalars().all()

    # 7. Fetch All Fuel Stocks
    stocks_q = select(FuelStock).options(selectinload(FuelStock.facility))
    if facility:
        stocks_q = stocks_q.where(FuelStock.facility_id == facility.id)
    stocks_list = (await db.execute(stocks_q.order_by(FuelStock.id.asc()))).scalars().all()

    # 8. Fetch Maintenance Records
    m_q = select(MaintenanceRecord).options(
        selectinload(MaintenanceRecord.user),
        selectinload(MaintenanceRecord.generator)
    ).where(
        MaintenanceRecord.created_at >= start_date,
        MaintenanceRecord.created_at <= end_date
    )
    if generator:
        m_q = m_q.where(MaintenanceRecord.generator_id == generator.id)
    elif facility:
        m_q = m_q.join(Generator).where(Generator.facility_id == facility.id)
    m_q = m_q.order_by(MaintenanceRecord.created_at.asc())
    maint_records = (await db.execute(m_q)).scalars().all()

    # 9. Fetch Faults
    f_q = select(Fault).options(selectinload(Fault.user), selectinload(Fault.generator)).where(
        Fault.created_at >= start_date,
        Fault.created_at <= end_date
    )
    if generator:
        f_q = f_q.where(Fault.generator_id == generator.id)
    elif facility:
        f_q = f_q.join(Generator).where(Generator.facility_id == facility.id)
    f_q = f_q.order_by(Fault.created_at.asc())
    faults = (await db.execute(f_q)).scalars().all()

    # Create Workbook
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # -------------------------------------------------------------
    # ВКЛАДКА 1: ЗАГАЛЬНИЙ ЗВІТ
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="Загальний звіт")
    ws1.views.sheetView[0].showGridLines = True

    # Title
    ws1.merge_cells("A1:F1")
    title_cell = ws1["A1"]
    title_cell.value = "SERVICE DESK (СЕРВІС ДЕСК) БЕНЗИНОВОГО ГЕНЕРАТОРА — ЗАГАЛЬНИЙ ЗВІТ"
    title_cell.font = TITLE_FONT
    title_cell.alignment = ALIGN_LEFT
    ws1.row_dimensions[1].height = 28

    # Metadata Block
    export_now_str = fmt_dt(now_facility())
    start_fmt = fmt_date(start_date)
    end_fmt = fmt_date(end_date)
    current_status_ua = GENERATOR_STATUS_UA.get(str(generator.status).upper(), str(generator.status)) if generator else "-"
    meta_info = [
        ("Об'єкт (локація):", facility_name, "Період звіту:", f"{start_fmt} — {end_fmt}"),
        ("Генератор:", f"{generator_name} ({gen_model})", "Дата експорту:", export_now_str),
        ("Серійний номер:", gen_serial, "Поточний статус:", current_status_ua),
        ("Поточні мотогодини:", f"{current_hours:.2f} год", "Запас палива в баку:", f"{generator.fuel_tank_level_l if generator else 0.0:.2f} л"),
    ]
    for row_idx, (k1, v1, k2, v2) in enumerate(meta_info, start=3):
        ws1.cell(row=row_idx, column=1, value=k1).font = BOLD_FONT
        ws1.cell(row=row_idx, column=2, value=v1).font = REGULAR_FONT
        ws1.cell(row=row_idx, column=4, value=k2).font = BOLD_FONT
        ws1.cell(row=row_idx, column=5, value=v2).font = REGULAR_FONT
        ws1.row_dimensions[row_idx].height = 20

    # Summary KPI Tables
    tot_hours_val = op_report.total_operating_hours_human or fmt_duration_human(None, op_report.total_operating_hours)
    tot_hours_exact = f"({fmt_duration(None, op_report.total_operating_hours)})"

    kpi_sections = [
        (
            "1. Експлуатація та напрацювання",
            SECTION_FILL,
            [
                ("Кількість робочих циклів (пусків)", op_report.total_runs_count, "шт"),
                ("Відпрацьовано мотогодин за період", tot_hours_val, tot_hours_exact),
                ("Розрахункова витрата бензину", round(op_report.total_calculated_consumption_l, 2), "л"),
                ("Поточні загальні мотогодини генератора", round(current_hours, 2), "год"),
            ]
        ),
        (
            "2. Паливо та запаси",
            PRIMARY_FILL,
            [
                ("Оприбутковано палива на склад", round(op_report.total_fuel_received_l, 2), "л"),
                ("Загальна вартість закупленого палива", round(op_report.total_fuel_receipts_cost_uah, 2), "грн"),
                ("Середня закупівельна ціна літра", round(op_report.avg_fuel_price_per_liter, 2), "грн/л"),
                ("Заправлено з складу в бак генератора", round(op_report.total_fuel_transferred_to_tank_l, 2), "л"),
                ("Поточний залишок на складі ГСМ", round(op_report.current_stock_balance_l, 2), "л"),
                ("Поточний рівень у паливному баку", round(op_report.current_tank_level_l, 2), "л"),
            ]
        ),
        (
            "3. Технічне обслуговування (ТО)",
            SECONDARY_FILL,
            [
                ("Проведено планових регламентних ТО", op_report.scheduled_maintenance_count, "шт"),
                ("Проведено проміжних (додаткових) ТО", op_report.intermediate_maintenance_count, "шт"),
                ("Загальні витрати на ТО та матеріали", round(op_report.total_maintenance_cost_uah, 2), "грн"),
                ("Нормативний інтервал ТО", round(interval_hours, 1), "год"),
                ("Останнє пройдене ТО на відмітці", round(last_maint_hours, 2), "год"),
                ("Наступне планове ТО на відмітці", round(next_due_hours, 2), "год"),
                ("Залишилося мотогодин до наступного ТО", round(hours_remaining, 2), "год"),
            ]
        ),
        (
            "4. Несправності та інциденти",
            SECTION_FILL,
            [
                ("Зареєстровано інцидентів / несправностей", op_report.total_faults_count, "шт"),
                ("Усунено та закрито інцидентів", op_report.resolved_faults_count, "шт"),
                ("Відкритих активних несправностей", op_report.open_faults_count, "шт"),
            ]
        )
    ]

    curr_row = 8
    for sec_title, fill_style, items in kpi_sections:
        ws1.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
        h_cell = ws1.cell(row=curr_row, column=1, value=sec_title)
        h_cell.fill = fill_style
        h_cell.font = SECTION_FONT
        h_cell.alignment = ALIGN_LEFT
        ws1.row_dimensions[curr_row].height = 24
        curr_row += 1

        for label, val, unit in items:
            c1 = ws1.cell(row=curr_row, column=1, value=label)
            c2 = ws1.cell(row=curr_row, column=2, value=val)
            c3 = ws1.cell(row=curr_row, column=3, value=unit)
            c1.font = REGULAR_FONT
            c2.font = BOLD_FONT
            c3.font = REGULAR_FONT
            c1.border = CELL_BORDER
            c2.border = CELL_BORDER
            c3.border = CELL_BORDER
            c1.alignment = ALIGN_LEFT
            c2.alignment = ALIGN_RIGHT
            c3.alignment = ALIGN_CENTER
            if unit in ("грн", "грн/л"):
                c2.number_format = "#,##0.00"
            elif isinstance(val, float):
                c2.number_format = "#,##0.00"
            elif isinstance(val, int):
                c2.number_format = "#,##0"
            ws1.row_dimensions[curr_row].height = 20
            curr_row += 1
        curr_row += 1

    auto_adjust_column_widths(ws1, min_width=15)

    # -------------------------------------------------------------
    # ВКЛАДКА 2: ЗВІТ ПО ДНЯХ (РОБОЧІ ЦИКЛИ ТА ОПЕРАЦІЇ)
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Звіт по днях")
    ws2.views.sheetView[0].showGridLines = True
    ws2.freeze_panes = "A2"

    headers2 = [
        "№", "Дата (ДД.ММ.РРРР)", "Час старту", "Час зупинки",
        "Тривалість (ГГ:ХХ:СС)", "Тривалість (год, хв)",
        "Початкові мотогодини", "Кінцеві мотогодини",
        "Початкове паливо в баку (л)", "Кінцеве паливо в баку (л)",
        "Витрата палива (л)", "Заправка в бак (л)", "Приход на склад (л)",
        "Оператор", "Статус", "Примітки"
    ]
    ws2.row_dimensions[1].height = 26
    for col_idx, h in enumerate(headers2, start=1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.fill = PRIMARY_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = CELL_BORDER

    # Map daily transfers and receipts by date
    transfers_by_date = {}
    for tr in transfers:
        d_str = fmt_date(tr.created_at)
        transfers_by_date[d_str] = transfers_by_date.get(d_str, 0.0) + (tr.liters or 0.0)

    receipts_by_date = {}
    for rc in receipts:
        d_str = fmt_date(rc.created_at)
        receipts_by_date[d_str] = receipts_by_date.get(d_str, 0.0) + (rc.liters or 0.0)

    tot_dur_hours = 0.0
    tot_cons_l = 0.0
    tot_fuel_added = 0.0
    tot_fuel_received = 0.0

    daily_rows = []
    seen_dates = set()

    for r in runs:
        date_str = fmt_date(r.start_time)
        seen_dates.add(date_str)
        start_t = fmt_time(r.start_time)
        end_t = fmt_time(r.end_time) if r.end_time else "-"
        dur_str = getattr(r, "duration_formatted", None) or fmt_duration(r.duration_seconds, r.duration_hours)
        dur_human = getattr(r, "duration_human", None) or fmt_duration_human(r.duration_seconds, r.duration_hours)
        consumption = round(r.calculated_consumption_l or 0.0, 2)
        fuel_added = round(transfers_by_date.pop(date_str, 0.0), 2)
        fuel_rec = round(receipts_by_date.pop(date_str, 0.0), 2)
        op_name = r.user.full_name if r.user else "Оператор"

        tot_dur_hours += (r.duration_hours or 0.0)
        tot_cons_l += consumption
        tot_fuel_added += fuel_added
        tot_fuel_received += fuel_rec

        daily_rows.append([
            date_str,
            start_t,
            end_t,
            dur_str,
            dur_human,
            round(r.start_hours, 3) if r.start_hours is not None else "-",
            round(r.end_hours, 3) if r.end_hours is not None else "-",
            round(r.start_fuel_level_l, 2) if r.start_fuel_level_l is not None else "-",
            round(r.end_fuel_level_l, 2) if r.end_fuel_level_l is not None else "-",
            consumption,
            fuel_added if fuel_added > 0 else "-",
            fuel_rec if fuel_rec > 0 else "-",
            op_name,
            RUN_STATUS_UA.get(str(r.status).upper(), str(r.status)),
            r.note or ""
        ])

    # Add any dates that had fuel movements (receipts or transfers) but no generator runs
    extra_dates = sorted(list(set(transfers_by_date.keys()) | set(receipts_by_date.keys())))
    for d_str in extra_dates:
        fuel_added = round(transfers_by_date.get(d_str, 0.0), 2)
        fuel_rec = round(receipts_by_date.get(d_str, 0.0), 2)
        tot_fuel_added += fuel_added
        tot_fuel_received += fuel_rec

        daily_rows.append([
            d_str,
            "-",
            "-",
            "-",
            "-",
            "-",
            "-",
            "-",
            "-",
            0.0,
            fuel_added if fuel_added > 0 else "-",
            fuel_rec if fuel_rec > 0 else "-",
            "Склад ГСМ",
            "Рух палива",
            "Оприбуткування / заправка без пуску генератора"
        ])

    for idx, vals in enumerate(daily_rows, start=1):
        r_row = idx + 1
        ws2.row_dimensions[r_row].height = 20
        fill = ZEBRA_FILL if idx % 2 == 0 else PatternFill(fill_type=None)

        # №
        c_num = ws2.cell(row=r_row, column=1, value=idx)
        c_num.font = REGULAR_FONT
        c_num.border = CELL_BORDER
        c_num.alignment = ALIGN_CENTER
        if fill.fill_type:
            c_num.fill = fill

        for c_offset, val in enumerate(vals, start=2):
            cell = ws2.cell(row=r_row, column=c_offset, value=sanitize_spreadsheet_cell(val))
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            if fill.fill_type:
                cell.fill = fill

            if c_offset in [2, 3, 4, 5, 6, 15]:
                cell.alignment = ALIGN_CENTER
            elif c_offset in [7, 8, 9, 10, 11, 12, 13]:
                cell.alignment = ALIGN_RIGHT
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = ALIGN_LEFT

    # Total row for Sheet 2
    last_r = len(daily_rows) + 2
    ws2.row_dimensions[last_r].height = 22
    for c_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(row=last_r, column=c_idx)
        cell.fill = TOTAL_FILL
        cell.border = TOTAL_BORDER
        cell.font = TOTAL_FONT

    ws2.cell(row=last_r, column=1, value="РАЗОМ").alignment = ALIGN_CENTER
    tot_dur_formatted = fmt_duration(None, tot_dur_hours)
    tot_dur_human = fmt_duration_human(None, tot_dur_hours)

    c_tot_exact = ws2.cell(row=last_r, column=5, value=tot_dur_formatted)
    c_tot_exact.alignment = ALIGN_CENTER

    c_tot_human = ws2.cell(row=last_r, column=6, value=tot_dur_human)
    c_tot_human.alignment = ALIGN_CENTER

    c_tot_cons = ws2.cell(row=last_r, column=11, value=round(tot_cons_l, 2))
    c_tot_cons.alignment = ALIGN_RIGHT
    c_tot_cons.number_format = "#,##0.00"
    c_tot_cons.number_format = "#,##0.00"

    if tot_fuel_added > 0:
        c_tot_add = ws2.cell(row=last_r, column=12, value=round(tot_fuel_added, 2))
        c_tot_add.alignment = ALIGN_RIGHT
        c_tot_add.number_format = "#,##0.00"

    if tot_fuel_received > 0:
        c_tot_rec = ws2.cell(row=last_r, column=13, value=round(tot_fuel_received, 2))
        c_tot_rec.alignment = ALIGN_RIGHT
        c_tot_rec.number_format = "#,##0.00"

    ws2.auto_filter.ref = f"A1:{get_column_letter(len(headers2))}{last_r - 1}"
    auto_adjust_column_widths(ws2, min_width=12)

    # -------------------------------------------------------------
    # ВКЛАДКА 3: РУХ ПАЛИВА ТА СКЛАД ГСМ
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Рух палива та Склад ГСМ")
    ws3.views.sheetView[0].showGridLines = True
    ws3.freeze_panes = "A3"

    # Section 1: Receipts
    ws3.merge_cells("A1:L1")
    s1_title = ws3["A1"]
    s1_title.value = "1. Оприбуткування палива на склад ГСМ (закупівля / чеки)"
    s1_title.font = SECTION_FONT
    s1_title.fill = PRIMARY_FILL
    s1_title.alignment = ALIGN_LEFT
    ws3.row_dimensions[1].height = 24

    headers3_rec = [
        "№", "Дата / Час", "Об'єкт (локація)", "Склад ГСМ", "Чек / Накладна №",
        "Водій / Постачальник", "Тип палива", "Об'єм (л)",
        "Сума (грн)", "Ціна за літр (грн/л)", "Хто прийняв", "Коментар"
    ]
    ws3.row_dimensions[2].height = 24
    for c_idx, h in enumerate(headers3_rec, start=1):
        cell = ws3.cell(row=2, column=c_idx, value=h)
        cell.fill = SECONDARY_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = CELL_BORDER

    tot_rec_l = 0.0
    tot_rec_uah = 0.0
    row_cur = 3
    for idx, rc in enumerate(receipts, start=1):
        tot_rec_l += rc.liters
        tot_rec_uah += rc.cost_total
        fac_rec_name = rc.stock.facility.name if rc.stock and rc.stock.facility else facility_name
        stock_name = rc.stock.name if rc.stock else "Склад ГСМ"
        rec_user = rc.user.full_name if rc.user else "Користувач"
        vals = [
            idx,
            fmt_dt(rc.created_at),
            fac_rec_name,
            stock_name,
            rc.receipt_number,
            rc.driver_name,
            rc.fuel_type,
            rc.liters,
            rc.cost_total,
            rc.price_per_liter,
            rec_user,
            rc.comment or ""
        ]
        ws3.row_dimensions[row_cur].height = 20
        fill = ZEBRA_FILL if idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, val in enumerate(vals, start=1):
            cell = ws3.cell(row=row_cur, column=c_idx, value=sanitize_spreadsheet_cell(val))
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            if fill.fill_type:
                cell.fill = fill
            if c_idx in [1, 2, 5, 7]:
                cell.alignment = ALIGN_CENTER
            elif c_idx in [8, 9, 10]:
                cell.alignment = ALIGN_RIGHT
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = ALIGN_LEFT
        row_cur += 1

    # Receipts Total Row
    ws3.row_dimensions[row_cur].height = 22
    for c_idx in range(1, len(headers3_rec) + 1):
        cell = ws3.cell(row=row_cur, column=c_idx)
        cell.fill = TOTAL_FILL
        cell.border = TOTAL_BORDER
        cell.font = TOTAL_FONT
    ws3.cell(row=row_cur, column=1, value="РАЗОМ").alignment = ALIGN_CENTER
    c_tot_l = ws3.cell(row=row_cur, column=8, value=round(tot_rec_l, 2))
    c_tot_l.alignment = ALIGN_RIGHT
    c_tot_l.number_format = "#,##0.00"
    c_tot_u = ws3.cell(row=row_cur, column=9, value=round(tot_rec_uah, 2))
    c_tot_u.alignment = ALIGN_RIGHT
    c_tot_u.number_format = "#,##0.00"
    if tot_rec_l > 0:
        c_avg = ws3.cell(row=row_cur, column=10, value=round(tot_rec_uah / tot_rec_l, 2))
        c_avg.alignment = ALIGN_RIGHT
        c_avg.number_format = "#,##0.00"
    row_cur += 3

    # Section 2: Transfers to Tank
    ws3.merge_cells(start_row=row_cur, start_column=1, end_row=row_cur, end_column=12)
    s2_title = ws3.cell(row=row_cur, column=1, value="2. Заправки зі складу ГСМ у бак генератора")
    s2_title.font = SECTION_FONT
    s2_title.fill = PRIMARY_FILL
    s2_title.alignment = ALIGN_LEFT
    ws3.row_dimensions[row_cur].height = 24
    row_cur += 1

    headers3_tr = [
        "№", "Дата / Час", "Об'єкт (локація)", "Звідки (Склад ГСМ)", "Куди (Генератор)",
        "Заправлено (л)", "Склад до (л)", "Склад після (л)",
        "Бак до (л)", "Бак після (л)", "Оператор", "Коментар"
    ]
    ws3.row_dimensions[row_cur].height = 24
    for c_idx, h in enumerate(headers3_tr, start=1):
        cell = ws3.cell(row=row_cur, column=c_idx, value=h)
        cell.fill = SECONDARY_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = CELL_BORDER
    row_cur += 1

    tot_tr_l = 0.0
    for idx, tr in enumerate(transfers, start=1):
        tot_tr_l += tr.liters
        fac_tr_name = tr.stock.facility.name if tr.stock and tr.stock.facility else facility_name
        stock_name = tr.stock.name if tr.stock else "Склад ГСМ"
        gen_tr_name = tr.generator.name if tr.generator else generator_name
        tr_user = tr.user.full_name if tr.user else "Оператор"
        vals = [
            idx,
            fmt_dt(tr.created_at),
            fac_tr_name,
            stock_name,
            gen_tr_name,
            tr.liters,
            tr.source_balance_before,
            tr.source_balance_after,
            tr.tank_balance_before,
            tr.tank_balance_after,
            tr_user,
            tr.comment or ""
        ]
        ws3.row_dimensions[row_cur].height = 20
        fill = ZEBRA_FILL if idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, val in enumerate(vals, start=1):
            cell = ws3.cell(row=row_cur, column=c_idx, value=sanitize_spreadsheet_cell(val))
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            if fill.fill_type:
                cell.fill = fill
            if c_idx in [1, 2]:
                cell.alignment = ALIGN_CENTER
            elif c_idx in [6, 7, 8, 9, 10]:
                cell.alignment = ALIGN_RIGHT
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = ALIGN_LEFT
        row_cur += 1

    # Transfers Total Row
    ws3.row_dimensions[row_cur].height = 22
    for c_idx in range(1, len(headers3_tr) + 1):
        cell = ws3.cell(row=row_cur, column=c_idx)
        cell.fill = TOTAL_FILL
        cell.border = TOTAL_BORDER
        cell.font = TOTAL_FONT
    ws3.cell(row=row_cur, column=1, value="РАЗОМ").alignment = ALIGN_CENTER
    c_tot_trl = ws3.cell(row=row_cur, column=6, value=round(tot_tr_l, 2))
    c_tot_trl.alignment = ALIGN_RIGHT
    c_tot_trl.number_format = "#,##0.00"
    row_cur += 3

    # Section 3: Current Balances of all Fuel Stocks
    ws3.merge_cells(start_row=row_cur, start_column=1, end_row=row_cur, end_column=6)
    s3_title = ws3.cell(row=row_cur, column=1, value="3. Поточні залишки на складах ГСМ")
    s3_title.font = SECTION_FONT
    s3_title.fill = PRIMARY_FILL
    s3_title.alignment = ALIGN_LEFT
    ws3.row_dimensions[row_cur].height = 24
    row_cur += 1

    headers3_stk = [
        "№", "Назва складу ГСМ", "Об'єкт (локація)", "Тип палива",
        "Поточний залишок (л)", "Дата останнього оновлення"
    ]
    ws3.row_dimensions[row_cur].height = 24
    for c_idx, h in enumerate(headers3_stk, start=1):
        cell = ws3.cell(row=row_cur, column=c_idx, value=h)
        cell.fill = SECONDARY_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = CELL_BORDER
    row_cur += 1

    tot_stk_l = 0.0
    for idx, stk in enumerate(stocks_list, start=1):
        tot_stk_l += stk.current_balance_l
        stk_fac = stk.facility.name if stk.facility else facility_name
        vals = [
            idx,
            stk.name,
            stk_fac,
            stk.fuel_type,
            stk.current_balance_l,
            fmt_dt(stk.updated_at)
        ]
        ws3.row_dimensions[row_cur].height = 20
        fill = ZEBRA_FILL if idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, val in enumerate(vals, start=1):
            cell = ws3.cell(row=row_cur, column=c_idx, value=sanitize_spreadsheet_cell(val))
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            if fill.fill_type:
                cell.fill = fill
            if c_idx in [1, 4, 6]:
                cell.alignment = ALIGN_CENTER
            elif c_idx == 5:
                cell.alignment = ALIGN_RIGHT
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = ALIGN_LEFT
        row_cur += 1

    # Stocks Total Row
    ws3.row_dimensions[row_cur].height = 22
    for c_idx in range(1, len(headers3_stk) + 1):
        cell = ws3.cell(row=row_cur, column=c_idx)
        cell.fill = TOTAL_FILL
        cell.border = TOTAL_BORDER
        cell.font = TOTAL_FONT
    ws3.cell(row=row_cur, column=1, value="РАЗОМ НА СКЛАДАХ").alignment = ALIGN_CENTER
    c_tot_stk = ws3.cell(row=row_cur, column=5, value=round(tot_stk_l, 2))
    c_tot_stk.alignment = ALIGN_RIGHT
    c_tot_stk.number_format = "#,##0.00"

    auto_adjust_column_widths(ws3, min_width=12)

    # -------------------------------------------------------------
    # ВКЛАДКА 4: ТЕХНІЧНЕ ОБСЛУГОВУВАННЯ (ТО)
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Технічне обслуговування")
    ws4.views.sheetView[0].showGridLines = True
    ws4.freeze_panes = "A2"

    headers4 = [
        "№", "Дата / Час", "Генератор", "Тип ТО",
        "Мотогодини при ТО", "Опис виконаних робіт",
        "Використані матеріали / запчастини", "Вартість (грн)",
        "Виконавець", "Примітки"
    ]
    ws4.row_dimensions[1].height = 26
    for c_idx, h in enumerate(headers4, start=1):
        cell = ws4.cell(row=1, column=c_idx, value=h)
        cell.fill = PRIMARY_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = CELL_BORDER

    tot_maint_cost = 0.0
    for idx, mr in enumerate(maint_records, start=1):
        r_row = idx + 1
        tot_maint_cost += mr.cost or 0.0
        m_gen = mr.generator.name if mr.generator else generator_name
        m_user = mr.user.full_name if mr.user else "Майстер"
        type_trans = "Регламентне ТО" if mr.maintenance_type == "SCHEDULED" else "Проміжне ТО"
        vals = [
            idx,
            fmt_dt(mr.created_at),
            m_gen,
            type_trans,
            mr.operating_hours,
            mr.work_description,
            mr.consumables_used or "-",
            mr.cost,
            m_user,
            mr.comment or ""
        ]
        ws4.row_dimensions[r_row].height = 20
        fill = ZEBRA_FILL if idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, val in enumerate(vals, start=1):
            cell = ws4.cell(row=r_row, column=c_idx, value=sanitize_spreadsheet_cell(val))
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            if fill.fill_type:
                cell.fill = fill
            if c_idx in [1, 2, 4]:
                cell.alignment = ALIGN_CENTER
            elif c_idx in [5, 8]:
                cell.alignment = ALIGN_RIGHT
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = ALIGN_LEFT

    # Total row Sheet 4
    last_r4 = len(maint_records) + 2
    ws4.row_dimensions[last_r4].height = 22
    for c_idx in range(1, len(headers4) + 1):
        cell = ws4.cell(row=last_r4, column=c_idx)
        cell.fill = TOTAL_FILL
        cell.border = TOTAL_BORDER
        cell.font = TOTAL_FONT
    ws4.cell(row=last_r4, column=1, value="РАЗОМ").alignment = ALIGN_CENTER
    c_m_tot = ws4.cell(row=last_r4, column=8, value=round(tot_maint_cost, 2))
    c_m_tot.alignment = ALIGN_RIGHT
    c_m_tot.number_format = "#,##0.00"

    ws4.auto_filter.ref = f"A1:{get_column_letter(len(headers4))}{last_r4 - 1}"
    auto_adjust_column_widths(ws4, min_width=12)

    # -------------------------------------------------------------
    # ВКЛАДКА 5: ЖУРНАЛ НЕСПРАВНОСТЕЙ ТА ІНЦИДЕНТІВ
    # -------------------------------------------------------------
    ws5 = wb.create_sheet(title="Журнал несправностей та інцидентів")
    ws5.views.sheetView[0].showGridLines = True
    ws5.freeze_panes = "A3"

    # Faults Section
    ws5.merge_cells("A1:I1")
    f_title = ws5["A1"]
    f_title.value = "Журнал несправностей та інцидентів"
    f_title.font = SECTION_FONT
    f_title.fill = PRIMARY_FILL
    f_title.alignment = ALIGN_LEFT
    ws5.row_dimensions[1].height = 24

    headers5_f = [
        "№", "Дата / Час", "Генератор", "Статус", "Вплив на роботу",
        "Опис несправності", "Дії з усунення", "Хто зафіксував", "Вирішено"
    ]
    ws5.row_dimensions[2].height = 24
    for c_idx, h in enumerate(headers5_f, start=1):
        cell = ws5.cell(row=2, column=c_idx, value=h)
        cell.fill = SECONDARY_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = CELL_BORDER

    row_f = 3
    for idx, fl in enumerate(faults, start=1):
        f_user = fl.user.full_name if fl.user else "Користувач"
        f_gen = fl.generator.name if fl.generator else generator_name
        resolved_t = fmt_dt(fl.resolved_at) if fl.resolved_at else "Не усунено"
        st_ua = FAULT_STATUS_UA.get(str(fl.status).upper(), str(fl.status))
        raw_imp = getattr(fl, "operational_impact", fl.priority)
        imp_ua = FAULT_PRIORITY_UA.get(str(raw_imp).upper(), str(raw_imp))
        vals = [
            idx,
            fmt_dt(fl.created_at),
            sanitize_spreadsheet_cell(f_gen),
            sanitize_spreadsheet_cell(st_ua),
            sanitize_spreadsheet_cell(imp_ua),
            sanitize_spreadsheet_cell(fl.description),
            sanitize_spreadsheet_cell(getattr(fl, "resolution_actions", fl.resolution_notes) or "-"),
            sanitize_spreadsheet_cell(f_user),
            resolved_t
        ]
        ws5.row_dimensions[row_f].height = 20
        fill = ZEBRA_FILL if idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, val in enumerate(vals, start=1):
            cell = ws5.cell(row=row_f, column=c_idx, value=val)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            if fill.fill_type:
                cell.fill = fill
            if c_idx in [1, 2, 4, 5, 9]:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT
        row_f += 1

    if row_f > 3:
        ws5.auto_filter.ref = f"A2:I{row_f - 1}"

    auto_adjust_column_widths(ws5, min_width=12)

    # Save workbook to memory
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


ADJUSTMENT_FIELD_UA = {
    "current_operating_hours": "Мотогодини",
    "fuel_tank_level_l": "Рівень у баку (л)",
    "status": "Статус генератора",
    "nominal_consumption_l_per_h": "Нормативний розхід (л/год)",
    "current_balance_l": "Залишок на складі (л)",
    "interval_hours": "Інтервал ТО (год)",
    "next_due_hours": "Наступне ТО (мотогодини)",
    "last_performed_hours": "Останнє ТО (мотогодини)",
    "name": "Назва",
    "address": "Адреса",
}


async def generate_audit_excel_report(
    db: AsyncSession,
    facility_id: Optional[int] = None,
    start_date: Optional[datetime.datetime] = None,
    end_date: Optional[datetime.datetime] = None,
    action: Optional[str] = None,
    entity_type: Optional[str] = None
) -> bytes:
    """
    Генерує спеціалізований аудиторський звіт у форматі Excel (.xlsx) для Головного Суперадміністратора.
    Містить 2 окремі аркуші відповідно до стандартів ДСТУ 4163:2020:
    1. Журнал дій аудиту (Повний хронологічний аудит операцій персоналу)
    2. Системні коригування (Реєстр коригувань лічильників мотогодин та палива з підставами)
    """
    now = utc_now()
    if not end_date:
        end_date = now

    # 1. Fetch Audit Logs
    audit_q = select(AuditLog).options(selectinload(AuditLog.user))
    if start_date:
        audit_q = audit_q.where(AuditLog.created_at >= start_date)
    if end_date:
        audit_q = audit_q.where(AuditLog.created_at <= end_date)
    if action:
        audit_q = audit_q.where(AuditLog.action == action)
    if entity_type:
        audit_q = audit_q.where(AuditLog.entity_type == entity_type)
    audit_q = audit_q.order_by(AuditLog.created_at.desc())
    audit_logs = (await db.execute(audit_q)).scalars().all()

    # 2. Fetch System Adjustments
    adj_q = select(SystemAdjustment).options(
        selectinload(SystemAdjustment.user),
        selectinload(SystemAdjustment.facility)
    )
    if start_date:
        adj_q = adj_q.where(SystemAdjustment.created_at >= start_date)
    if end_date:
        adj_q = adj_q.where(SystemAdjustment.created_at <= end_date)
    if facility_id:
        adj_q = adj_q.where(SystemAdjustment.facility_id == facility_id)
    adj_q = adj_q.order_by(SystemAdjustment.created_at.desc())
    adjustments = (await db.execute(adj_q)).scalars().all()

    # Create Workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    start_fmt = fmt_date_ua(start_date) if start_date else "Початок ведення журналу"
    end_fmt = fmt_date_ua(end_date)
    export_now_str = fmt_dt(now_facility())

    # -------------------------------------------------------------
    # ВКЛАДКА 1: ЖУРНАЛ ДІЙ АУДИТУ
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="Журнал дій аудиту")
    ws1.views.sheetView[0].showGridLines = True
    ws1.freeze_panes = "A5"

    # Title
    ws1.merge_cells("A1:H1")
    t1 = ws1["A1"]
    t1.value = "SERVICE DESK SD-GEN016 — ЖУРНАЛ ОПЕРАЦІЙНОГО АУДИТУ"
    t1.font = TITLE_FONT
    t1.alignment = ALIGN_LEFT
    ws1.row_dimensions[1].height = 28

    # Metadata
    ws1.cell(row=2, column=1, value="Період звіту:").font = BOLD_FONT
    ws1.cell(row=2, column=2, value=f"{start_fmt} — {end_fmt}").font = REGULAR_FONT
    ws1.cell(row=2, column=4, value="Дата вивантаження:").font = BOLD_FONT
    ws1.cell(row=2, column=5, value=export_now_str).font = REGULAR_FONT

    ws1.cell(row=3, column=1, value="Рівень доступу:").font = BOLD_FONT
    ws1.cell(row=3, column=2, value="Головний адміністратор (Суперадмін)").font = REGULAR_FONT
    ws1.cell(row=3, column=4, value="Всього записів аудиту:").font = BOLD_FONT
    ws1.cell(row=3, column=5, value=len(audit_logs)).font = REGULAR_FONT
    ws1.row_dimensions[2].height = 20
    ws1.row_dimensions[3].height = 20

    # Headers
    headers1 = [
        "№", "Дата / Час", "Користувач", "Дія", "Сутність", "ID сутності", "Деталі операції", "IP-адреса"
    ]
    ws1.row_dimensions[4].height = 24
    for c_idx, h in enumerate(headers1, start=1):
        cell = ws1.cell(row=4, column=c_idx, value=h)
        cell.fill = SECONDARY_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = CELL_BORDER

    row1 = 5
    for idx, al in enumerate(audit_logs, start=1):
        a_user = al.user.full_name if al.user else (al.user.login if al.user else "Система")
        details_str = format_audit_details_ua(al.details_json, str(al.action))
        act_ua = AUDIT_ACTION_UA.get(str(al.action), str(al.action))
        ent_ua = AUDIT_ENTITY_UA.get(str(al.entity_type), str(al.entity_type))
        vals = [
            idx,
            fmt_dt(al.created_at),
            sanitize_spreadsheet_cell(a_user),
            sanitize_spreadsheet_cell(act_ua),
            sanitize_spreadsheet_cell(ent_ua),
            sanitize_spreadsheet_cell(str(al.entity_id or "-")),
            sanitize_spreadsheet_cell(details_str),
            sanitize_spreadsheet_cell(al.ip_address or "-"),
        ]
        ws1.row_dimensions[row1].height = 20
        fill = ZEBRA_FILL if idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, val in enumerate(vals, start=1):
            cell = ws1.cell(row=row1, column=c_idx, value=val)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            if fill.fill_type:
                cell.fill = fill
            if c_idx in [1, 2, 4, 5, 6, 8]:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT
        row1 += 1

    if row1 > 5:
        ws1.auto_filter.ref = f"A4:H{row1 - 1}"
    auto_adjust_column_widths(ws1, min_width=12)

    # -------------------------------------------------------------
    # ВКЛАДКА 2: СИСТЕМНІ КОРИГУВАННЯ
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Системні коригування")
    ws2.views.sheetView[0].showGridLines = True
    ws2.freeze_panes = "A5"

    # Title
    ws2.merge_cells("A1:J1")
    t2 = ws2["A1"]
    t2.value = "SERVICE DESK SD-GEN016 — РЕЄСТР СИСТЕМНИХ КОРИГУВАНЬ ЛІЧИЛЬНИКІВ"
    t2.font = TITLE_FONT
    t2.alignment = ALIGN_LEFT
    ws2.row_dimensions[1].height = 28

    # Metadata
    ws2.cell(row=2, column=1, value="Період звіту:").font = BOLD_FONT
    ws2.cell(row=2, column=2, value=f"{start_fmt} — {end_fmt}").font = REGULAR_FONT
    ws2.cell(row=2, column=4, value="Дата вивантаження:").font = BOLD_FONT
    ws2.cell(row=2, column=5, value=export_now_str).font = REGULAR_FONT

    ws2.cell(row=3, column=1, value="Підстава операцій:").font = BOLD_FONT
    ws2.cell(row=3, column=2, value="Акти інвентаризації та звірки показників").font = REGULAR_FONT
    ws2.cell(row=3, column=4, value="Всього коригувань:").font = BOLD_FONT
    ws2.cell(row=3, column=5, value=len(adjustments)).font = REGULAR_FONT
    ws2.row_dimensions[2].height = 20
    ws2.row_dimensions[3].height = 20

    # Headers
    headers2 = [
        "№", "Дата / Час", "Адміністратор", "Об'єкт (локація)", "Цільова сутність",
        "Тип сутності", "Поле коригування", "Старе значення", "Нове значення", "Причина (Акт інвентаризації)"
    ]
    ws2.row_dimensions[4].height = 24
    for c_idx, h in enumerate(headers2, start=1):
        cell = ws2.cell(row=4, column=c_idx, value=h)
        cell.fill = SECTION_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = CELL_BORDER

    row2 = 5
    for idx, adj in enumerate(adjustments, start=1):
        adj_user = adj.user.full_name if adj.user else (adj.user.login if adj.user else "Система")
        fac_name = adj.facility.name if adj.facility else (f"Об'єкт #{adj.facility_id}" if adj.facility_id else "-")
        ent_ua = AUDIT_ENTITY_UA.get(str(adj.entity_type), str(adj.entity_type))
        field_ua = ADJUSTMENT_FIELD_UA.get(str(adj.field_name), str(adj.field_name))
        target_disp = adj.target_name or f"{ent_ua} #{adj.entity_id}"
        vals = [
            idx,
            fmt_dt(adj.created_at),
            sanitize_spreadsheet_cell(adj_user),
            sanitize_spreadsheet_cell(fac_name),
            sanitize_spreadsheet_cell(target_disp),
            sanitize_spreadsheet_cell(ent_ua),
            sanitize_spreadsheet_cell(field_ua),
            sanitize_spreadsheet_cell(adj.old_value),
            sanitize_spreadsheet_cell(adj.new_value),
            sanitize_spreadsheet_cell(adj.reason),
        ]
        ws2.row_dimensions[row2].height = 20
        fill = ZEBRA_FILL if idx % 2 == 0 else PatternFill(fill_type=None)
        for c_idx, val in enumerate(vals, start=1):
            cell = ws2.cell(row=row2, column=c_idx, value=val)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            if fill.fill_type:
                cell.fill = fill
            if c_idx in [1, 2, 6, 8, 9]:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT
        row2 += 1

    if row2 > 5:
        ws2.auto_filter.ref = f"A4:J{row2 - 1}"
    auto_adjust_column_widths(ws2, min_width=12)

    # Save workbook to memory
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

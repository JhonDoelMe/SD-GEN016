import io
import openpyxl
import pytest


@pytest.mark.asyncio
async def test_system_adjustments_and_audit(client, superadmin_auth, operator_auth):
    # Setup generator
    gen_res = await client.post(
        "/api/v1/generator/wizard",
        headers=superadmin_auth,
        json={
            "name": "Генератор Аудиту",
            "model": "A-1",
            "manufacturer": "PG",
            "serial_number": "SN-AUDIT-01",
            "rated_power_kw": 5.0,
            "tank_capacity_l": 25.0,
            "fuel_type": "А-95",
            "nominal_consumption_l_per_h": 2.0,
            "initial_operating_hours": 50.0,
            "initial_fuel_tank_level_l": 10.0,
            "work_schedule_start": "00:00",
            "work_schedule_end": "23:59",
            "maintenance_interval_hours": 300.0,
            "timezone": "Europe/Kyiv"
        }
    )
    assert gen_res.status_code == 200
    gen_id = gen_res.json()["id"]

    # 1. Operator attempts system adjustment -> FORBIDDEN (403)
    op_adj = await client.post(
        "/api/v1/adjustments",
        headers=operator_auth,
        json={
            "entity_type": "Generator",
            "entity_id": gen_id,
            "field_name": "current_operating_hours",
            "new_value": "45.0",
            "reason": "Спроба оператора зменшити лічильник"
        }
    )
    assert op_adj.status_code == 403
    assert "виключно головному системному адміністратору" in op_adj.json()["detail"]

    # 2. SuperAdmin attempts adjustment without valid reason -> BAD REQUEST (400 or 422)
    no_reason = await client.post(
        "/api/v1/adjustments",
        headers=superadmin_auth,
        json={
            "entity_type": "Generator",
            "entity_id": gen_id,
            "field_name": "current_operating_hours",
            "new_value": "55.0",
            "reason": " "
        }
    )
    assert no_reason.status_code in (400, 422)

    # 3. SuperAdmin performs valid adjustment of operating hours
    valid_adj = await client.post(
        "/api/v1/adjustments",
        headers=superadmin_auth,
        json={
            "entity_type": "Generator",
            "entity_id": gen_id,
            "field_name": "current_operating_hours",
            "new_value": "55.0",
            "reason": "Коригування за актом звірки механічного лічильника №14"
        }
    )
    assert valid_adj.status_code == 201
    adj_data = valid_adj.json()
    assert adj_data["old_value"] == "50.0"
    assert adj_data["new_value"] == "55.0"
    assert "№14" in adj_data["reason"]

    # Verify generator current hours updated
    gen_check = await client.get("/api/v1/generator", headers=superadmin_auth)
    assert gen_check.json()["current_operating_hours"] == 55.0

    # 4. Create a FuelStock and perform warehouse adjustment
    stock_res = await client.post(
        "/api/v1/fuel/stocks",
        headers=superadmin_auth,
        json={
            "facility_id": 1,
            "name": "Основний резервуар ДП",
            "fuel_type": "ДП",
            "capacity_l": 1000.0,
            "initial_balance_l": 200.0,
            "min_balance_warning_l": 50.0,
            "min_balance_critical_l": 20.0
        }
    )
    assert stock_res.status_code == 201
    stock_id = stock_res.json()["id"]

    # Adjust fuel stock balance
    stock_adj = await client.post(
        "/api/v1/adjustments",
        headers=superadmin_auth,
        json={
            "facility_id": 1,
            "entity_type": "FuelStock",
            "entity_id": stock_id,
            "field_name": "current_balance_l",
            "new_value": "350.5",
            "reason": "Акт інвентаризації резервуара №3 від 20.10.2026"
        }
    )
    assert stock_adj.status_code == 201
    stock_adj_data = stock_adj.json()
    assert stock_adj_data["old_value"] == "200.0"
    assert stock_adj_data["new_value"] == "350.5"
    assert stock_adj_data["facility_id"] == 1
    assert "Основний резервуар ДП" in stock_adj_data["target_name"]

    # Verify fuel stock balance was actually updated in database
    stock_check = await client.get("/api/v1/fuel/stocks", headers=superadmin_auth)
    assert stock_check.status_code == 200
    stocks = stock_check.json()
    updated_stock = next(s for s in stocks if s["id"] == stock_id)
    assert updated_stock["current_balance_l"] == 350.5

    # 5. List adjustments with facility filter
    adj_list = await client.get("/api/v1/adjustments?facility_id=1", headers=superadmin_auth)
    assert adj_list.status_code == 200
    items = adj_list.json()
    assert len(items) >= 2
    assert any(i["entity_type"] == "FuelStock" and i["entity_id"] == stock_id for i in items)

    # 6. Check audit log contains both adjustments and previous operations
    audit_res = await client.get("/api/v1/audit", headers=superadmin_auth)
    assert audit_res.status_code == 200
    logs = audit_res.json()
    assert len(logs) > 0
    actions = [l["action"] for l in logs]
    assert "SYSTEM_ADJUSTMENT" in actions
    assert "GENERATOR_WIZARD_CONFIGURED" in actions


@pytest.mark.asyncio
async def test_adjustments_negative_fuel_and_generator_lookup(client, superadmin_auth):
    # Fetch existing generator via /generator/list
    list_res = await client.get("/api/v1/generator/list", headers=superadmin_auth)
    assert list_res.status_code == 200
    gens = list_res.json()
    assert len(gens) > 0
    gen_id = gens[0]["id"]

    # SuperAdmin adjusts fuel level to negative (e.g. -5.5 l to record deficit)
    adj_res = await client.post(
        "/api/v1/adjustments",
        headers=superadmin_auth,
        json={
            "entity_type": "Generator",
            "entity_id": gen_id,
            "field_name": "fuel_tank_level_l",
            "new_value": "-5.5",
            "reason": "Фіксація реального дефіциту палива після вичерпання бака"
        }
    )
    assert adj_res.status_code == 201
    assert adj_res.json()["new_value"] == "-5.5"

    # Verify generator fuel level in DB is now -5.5
    gen_check = await client.get(f"/api/v1/generator?generator_id={gen_id}", headers=superadmin_auth)
    assert gen_check.status_code == 200
    assert gen_check.json()["fuel_tank_level_l"] == -5.5


@pytest.mark.asyncio
async def test_audit_excel_export_and_rbac(client, superadmin_auth, admin_auth, operator_auth, viewer_auth):
    # 1. Operator attempts export -> 403 Forbidden
    op_res = await client.get("/api/v1/audit/export-excel", headers=operator_auth)
    assert op_res.status_code == 403
    assert "виключно головному системному адміністратору" in op_res.json()["detail"]

    # 2. Viewer attempts export -> 403 Forbidden
    view_res = await client.get("/api/v1/audit/export-excel", headers=viewer_auth)
    assert view_res.status_code == 403

    # 3. Regular admin attempts export -> 403 Forbidden (Audit excel export is strictly SuperAdmin)
    adm_res = await client.get("/api/v1/audit/export-excel", headers=admin_auth)
    assert adm_res.status_code == 403
    assert "виключно головному системному адміністратору" in adm_res.json()["detail"]

    # 4. SuperAdmin performs export with period filter
    sa_res = await client.get(
        "/api/v1/audit/export-excel?start_date=2020-01-01&end_date=2030-12-31",
        headers=superadmin_auth
    )
    assert sa_res.status_code == 200
    assert "spreadsheetml.sheet" in sa_res.headers["content-type"]
    assert "Audit_Report_" in sa_res.headers["content-disposition"]
    assert len(sa_res.content) > 1000

    # 5. Parse and inspect sheets
    wb = openpyxl.load_workbook(io.BytesIO(sa_res.content))
    assert "Журнал дій аудиту" in wb.sheetnames
    assert "Системні коригування" in wb.sheetnames

    # Check Sheet 1 contents
    ws1 = wb["Журнал дій аудиту"]
    assert ws1.freeze_panes == "A5"
    ws1_text = " ".join(str(cell.value) for row in ws1.iter_rows() for cell in row if cell.value)
    assert "ЖУРНАЛ ОПЕРАЦІЙНОГО АУДИТУ" in ws1_text

    # Check Sheet 2 contents
    ws2 = wb["Системні коригування"]
    assert ws2.freeze_panes == "A5"
    ws2_text = " ".join(str(cell.value) for row in ws2.iter_rows() for cell in row if cell.value)
    assert "РЕЄСТР СИСТЕМНИХ КОРИГУВАНЬ" in ws2_text

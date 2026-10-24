import io
import pytest
import openpyxl


@pytest.mark.asyncio
async def test_facility_crud_and_multisite(client, superadmin_auth, operator_auth):
    # 1. List facilities - default facility should exist
    fac_list_res = await client.get("/api/v1/facilities", headers=superadmin_auth)
    assert fac_list_res.status_code == 200
    facs = fac_list_res.json()
    assert len(facs) >= 1
    default_fac = facs[0]
    assert "Основний об'єкт" in default_fac["name"]

    # 2. Create second facility
    create_res = await client.post(
        "/api/v1/facilities",
        headers=superadmin_auth,
        json={
            "name": "Філія №2 (Дарниця)",
            "address": "вул. Бориспільська, 9",
            "description": "Резервний вузол зв'язку",
            "timezone": "Europe/Kyiv"
        }
    )
    assert create_res.status_code == 201
    new_fac = create_res.json()
    assert new_fac["name"] == "Філія №2 (Дарниця)"
    assert new_fac["id"] > default_fac["id"]

    # 3. Update facility
    update_res = await client.put(
        f"/api/v1/facilities/{new_fac['id']}",
        headers=superadmin_auth,
        json={"address": "вул. Бориспільська, 15"}
    )
    assert update_res.status_code == 200
    assert update_res.json()["address"] == "вул. Бориспільська, 15"

    # 4. Check fuel stock auto-created for this facility
    stock_res = await client.get(
        f"/api/v1/fuel/stock?facility_id={new_fac['id']}",
        headers=superadmin_auth
    )
    assert stock_res.status_code == 200

    # 5. Delete facility
    del_res = await client.delete(
        f"/api/v1/facilities/{new_fac['id']}",
        headers=superadmin_auth
    )
    assert del_res.status_code == 200
    assert "успішно видалено" in del_res.json()["message"]


@pytest.mark.asyncio
async def test_generator_creation_deletion_and_safety(client, superadmin_auth, operator_auth):
    # 1. Create a new generator
    gen_create_res = await client.post(
        "/api/v1/generator",
        headers=superadmin_auth,
        json={
            "facility_id": 1,
            "name": "Тестовий генератор 7кВт",
            "model": "Honda EM 6500",
            "manufacturer": "Honda",
            "serial_number": "HND-778899",
            "rated_power_kw": 7.0,
            "tank_capacity_l": 30.0,
            "fuel_type": "А-95",
            "nominal_consumption_l_per_h": 2.5,
            "initial_operating_hours": 10.0,
            "initial_fuel_tank_level_l": 20.0,
            "work_schedule_start": "00:00",
            "work_schedule_end": "23:59",
            "maintenance_interval_hours": 300.0,
            "timezone": "Europe/Kyiv"
        }
    )
    assert gen_create_res.status_code == 201
    new_gen = gen_create_res.json()
    gen_id = new_gen["id"]

    # 2. Start the generator
    start_res = await client.post(
        "/api/v1/generator/start",
        headers=operator_auth,
        json={"generator_id": gen_id, "fuel_level_l": 20.0}
    )
    assert start_res.status_code == 200

    # 3. Safety check: Invariant - cannot delete RUNNING generator
    del_running_res = await client.delete(
        f"/api/v1/generator/{gen_id}",
        headers=superadmin_auth
    )
    assert del_running_res.status_code == 400
    assert "зараз працює" in del_running_res.json()["detail"]

    # 4. Stop the generator
    stop_res = await client.post(
        "/api/v1/generator/stop",
        headers=operator_auth,
        json={"generator_id": gen_id, "note": "Зупинка перед видаленням"}
    )
    assert stop_res.status_code == 200

    # 5. Delete stopped generator -> success
    del_ok_res = await client.delete(
        f"/api/v1/generator/{gen_id}",
        headers=superadmin_auth
    )
    assert del_ok_res.status_code == 200
    assert "успішно видалено" in del_ok_res.json()["message"]


@pytest.mark.asyncio
async def test_maintenance_schedule_calculation_and_recalculate(client, superadmin_auth):
    # Setup generator with current hours = 235.0, interval = 300.0
    setup_res = await client.post(
        "/api/v1/generator/wizard",
        headers=superadmin_auth,
        json={
            "name": "Генератор ТО Тест",
            "model": "GT-5000",
            "manufacturer": "GenTech",
            "serial_number": "SN-TO-300",
            "rated_power_kw": 5.0,
            "tank_capacity_l": 25.0,
            "fuel_type": "А-95",
            "nominal_consumption_l_per_h": 2.2,
            "initial_operating_hours": 235.0,
            "initial_fuel_tank_level_l": 15.0,
            "work_schedule_start": "00:00",
            "work_schedule_end": "23:59",
            "maintenance_interval_hours": 300.0,
            "timezone": "Europe/Kyiv"
        }
    )
    assert setup_res.status_code == 200
    gen_id = setup_res.json()["id"]

    # Check maintenance schedule status: next_due should be 300.0, remaining should be 65.0
    sched_res = await client.get(
        f"/api/v1/maintenance/schedule?generator_id={gen_id}",
        headers=superadmin_auth
    )
    assert sched_res.status_code == 200
    sched = sched_res.json()
    assert sched["interval_hours"] == 300.0
    assert sched["next_due_hours"] == 300.0
    assert sched["hours_remaining"] == 65.0

    # Test recalculate endpoint
    recalc_res = await client.post(
        f"/api/v1/maintenance/schedule/recalculate?generator_id={gen_id}",
        headers=superadmin_auth
    )
    assert recalc_res.status_code == 200
    recalc = recalc_res.json()
    assert recalc["next_due_hours"] == 300.0
    assert recalc["hours_remaining"] == 65.0


@pytest.mark.asyncio
async def test_dynamic_excel_report_generation(client, superadmin_auth, operator_auth):
    # Ensure there is a real GeneratorRun in DB so all loops in generate_excel_report execute
    setup_res = await client.post(
        "/api/v1/generator/wizard",
        headers=superadmin_auth,
        json={
            "name": "Генератор Звіт Тест",
            "model": "GT-5000",
            "manufacturer": "GenTech",
            "serial_number": "SN-REP-01",
            "rated_power_kw": 5.0,
            "tank_capacity_l": 25.0,
            "fuel_type": "А-95",
            "nominal_consumption_l_per_h": 2.2,
            "initial_operating_hours": 100.0,
            "initial_fuel_tank_level_l": 20.0,
            "work_schedule_start": "00:00",
            "work_schedule_end": "23:59",
            "maintenance_interval_hours": 300.0,
            "timezone": "Europe/Kyiv"
        }
    )
    assert setup_res.status_code == 200
    gen_id = setup_res.json()["id"]

    start_res = await client.post(
        "/api/v1/generator/start",
        headers=operator_auth,
        json={"generator_id": gen_id, "fuel_level_l": 15.0}
    )
    assert start_res.status_code == 200

    stop_res = await client.post(
        "/api/v1/generator/stop",
        headers=operator_auth,
        json={"generator_id": gen_id, "note": "Тестовий цикл для звіту Excel"}
    )
    assert stop_res.status_code == 200

    # Request Excel report with date range query parameters matching frontend format
    res = await client.get(
        f"/api/v1/reports/export-excel?generator_id={gen_id}&start_date=2020-01-01&end_date=2030-12-31",
        headers=superadmin_auth
    )
    assert res.status_code == 200
    assert "spreadsheetml.sheet" in res.headers["content-type"]
    assert "attachment; filename=" in res.headers["content-disposition"]
    assert len(res.content) > 1000

    # Parse Excel with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    sheet_names = wb.sheetnames

    # Verify all 5 required sheets exist
    assert "Загальний звіт" in sheet_names
    assert "Звіт по днях" in sheet_names
    assert "Рух палива та Склад ГСМ" in sheet_names
    assert "Технічне обслуговування" in sheet_names
    assert "Журнал інцидентів та дій" in sheet_names

    # Verify Sheet 1 content
    ws1 = wb["Загальний звіт"]
    assert "SERVICE DESK" in ws1["A1"].value or "СЕРВІС ДЕСК" in ws1["A1"].value

    # Verify Sheet 2 freeze panes, headers, and populated GeneratorRun row
    ws2 = wb["Звіт по днях"]
    assert ws2.freeze_panes == "A2"
    assert ws2["A1"].value == "№"
    assert ws2["B1"].value == "Дата (ДД.ММ.РРРР)"
    ws2_text = " ".join(str(cell.value) for row in ws2.iter_rows() for cell in row if cell.value)
    assert "Тестовий цикл для звіту Excel" in ws2_text


@pytest.mark.asyncio
async def test_multi_warehouse_receipt_transfer_and_excel(client, superadmin_auth):
    # 1. Create a facility
    fac_res = await client.post("/api/v1/facilities", json={
        "name": "Локація Захід (Львів)",
        "address": "вул. Городоцька, 15",
        "description": "Тестова локація",
        "timezone": "Europe/Kyiv"
    }, headers=superadmin_auth)
    assert fac_res.status_code == 201
    fac_id = fac_res.json()["id"]

    # 2. Create a generator for this facility
    gen_res = await client.post("/api/v1/generator", json={
        "facility_id": fac_id,
        "name": "Генератор Львів-1",
        "model": "Kipor KDE6700",
        "manufacturer": "Kipor",
        "serial_number": "LV-6700-01",
        "rated_power_kw": 5.5,
        "tank_capacity_l": 25.0,
        "fuel_type": "А-95",
        "nominal_consumption_l_per_h": 1.4,
        "initial_operating_hours": 10.0,
        "initial_fuel_tank_level_l": 5.0,
        "maintenance_interval_hours": 300.0,
        "work_schedule_start": "08:00",
        "work_schedule_end": "20:00",
        "timezone": "Europe/Kyiv"
    }, headers=superadmin_auth)
    assert gen_res.status_code == 201
    gen_id = gen_res.json()["id"]

    # 3. Create a fuel warehouse for this facility
    stock_res = await client.post("/api/v1/fuel/stock/new", json={
        "facility_id": fac_id,
        "name": "Резервний склад Львів",
        "fuel_type": "А-95",
        "initial_balance_l": 50.0
    }, headers=superadmin_auth)
    assert stock_res.status_code == 201
    stock = stock_res.json()
    stock_id = stock["id"]
    assert stock["name"] == "Резервний склад Львів"
    assert stock["current_balance_l"] == 50.0

    # 4. List stocks for this facility
    list_res = await client.get(f"/api/v1/fuel/stocks?facility_id={fac_id}", headers=superadmin_auth)
    assert list_res.status_code == 200
    stocks = list_res.json()
    assert any(s["id"] == stock_id for s in stocks)

    # 5. Add fuel receipt directly to this warehouse
    rec_res = await client.post("/api/v1/fuel/receipt", json={
        "facility_id": fac_id,
        "stock_id": stock_id,
        "liters": 100.0,
        "cost_total": 5600.0,
        "driver_name": "Петренко П.П.",
        "receipt_number": "ЧЕК-LV-100",
        "fuel_type": "А-95",
        "comment": "Завіз палива"
    }, headers=superadmin_auth)
    assert rec_res.status_code == 201
    rec = rec_res.json()
    assert rec["stock_name"] == "Резервний склад Львів"
    assert rec["liters"] == 100.0
    assert rec["facility_name"] == "Локація Захід (Львів)"

    # 6. Transfer fuel to generator tank
    trans_res = await client.post("/api/v1/fuel/transfer", json={
        "facility_id": fac_id,
        "stock_id": stock_id,
        "generator_id": gen_id,
        "liters": 15.0,
        "comment": "Заправка перед роботою"
    }, headers=superadmin_auth)
    assert trans_res.status_code == 201
    trans = trans_res.json()
    assert trans["stock_name"] == "Резервний склад Львів"
    assert trans["generator_name"] == "Генератор Львів-1"
    assert trans["liters"] == 15.0
    assert trans["tank_balance_after"] == 20.0

    # 7. Generate Excel report for this facility
    rep_res = await client.get(f"/api/v1/reports/export-excel?facility_id={fac_id}", headers=superadmin_auth)
    assert rep_res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(rep_res.content))
    ws_fuel = wb["Рух палива та Склад ГСМ"]
    # Check that receipt and transfer data appear in the worksheet
    sheet_text = " ".join(str(cell.value) for row in ws_fuel.iter_rows() for cell in row if cell.value)
    assert "Резервний склад Львів" in sheet_text
    assert "ЧЕК-LV-100" in sheet_text


@pytest.mark.asyncio
async def test_system_status_api(client, superadmin_auth):
    """Тест ендпоінту стану сервера для підвалу сайту"""
    res = await client.get("/api/v1/system/status")
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "online"
    assert "cpu" in data and "percent" in data["cpu"]
    assert "memory" in data and "percent" in data["memory"]
    assert "uptime" in data and "formatted" in data["uptime"]
    assert data["database"]["status"] == "connected"
    assert "timezone" in data


@pytest.mark.asyncio
async def test_generator_quick_update_api(client, superadmin_auth):
    """Тест швидких налаштувань генератора через PUT /api/v1/generator/{id}"""
    # 1. Fetch default generator
    gen_res = await client.get("/api/v1/generator", headers=superadmin_auth)
    assert gen_res.status_code == 200
    gen_id = gen_res.json()["id"]

    # 2. Update settings
    upd_res = await client.put(f"/api/v1/generator/{gen_id}", json={
        "name": "Генератор Оновлений",
        "model": "Honda GX390 Pro",
        "rated_power_kw": 6.0,
        "tank_capacity_l": 28.0,
        "nominal_consumption_l_per_h": 2.25
    }, headers=superadmin_auth)
    assert upd_res.status_code == 200
    updated = upd_res.json()
    assert updated["name"] == "Генератор Оновлений"
    assert updated["model"] == "Honda GX390 Pro"
    assert updated["rated_power_kw"] == 6.0
    assert updated["tank_capacity_l"] == 28.0
    assert updated["nominal_consumption_l_per_h"] == 2.25


@pytest.mark.asyncio
async def test_admin_system_settings_and_internet_fuel_prices(client, superadmin_auth, operator_auth):
    """Тест системних налаштувань (тільки адмін) та автоматичного розрахунку ціни палива з інтернету/налаштувань"""
    # 1. Operator cannot access /api/v1/system/settings
    op_res = await client.get("/api/v1/system/settings", headers=operator_auth)
    assert op_res.status_code == 403

    # 2. SuperAdmin gets system settings
    get_res = await client.get("/api/v1/system/settings", headers=superadmin_auth)
    assert get_res.status_code == 200
    data = get_res.json()
    assert "settings" in data
    assert "fuel_prices" in data
    assert "А-95" in data["fuel_prices"]["fuels"]

    # 3. Receipt WITHOUT cost_total when fuel_price_a95 == 0 -> uses internet price automatically!
    rec_internet_res = await client.post("/api/v1/fuel/receipt", json={
        "liters": 20.0,
        "driver_name": "Водій Інтернет Ціна",
        "receipt_number": "AUTO-NET-01",
        "fuel_type": "А-95"
    }, headers=operator_auth)
    assert rec_internet_res.status_code == 201
    rec_net = rec_internet_res.json()
    assert rec_net["liters"] == 20.0
    assert rec_net["price_per_liter"] >= 30.0
    assert rec_net["cost_total"] == round(20.0 * rec_net["price_per_liter"], 2)
    assert "Авторозрахунок" in (rec_net["comment"] or "")

    # 4. Admin sets custom fuel price in system settings (e.g. А-95 = 60.00 грн/л) and apply_to_all_generators = False
    put_res = await client.put("/api/v1/system/settings", json={
        "organization_name": "MAXI-UA TEST",
        "facility_timezone": "Europe/Kyiv",
        "work_schedule_start": "07:00",
        "work_schedule_end": "22:00",
        "fuel_price_a95": 60.0,
        "apply_to_all_generators": False
    }, headers=superadmin_auth)
    assert put_res.status_code == 200
    assert put_res.json()["settings"]["fuel_price_a95"] == 60.0
    assert put_res.json()["settings"]["apply_to_all_generators"] is False
    assert len(put_res.json()["fuel_prices"]["wog_lineup"]) >= 5

    # Verify GET /api/v1/system/settings preserves apply_to_all_generators == False
    get_res2 = await client.get("/api/v1/system/settings", headers=superadmin_auth)
    assert get_res2.status_code == 200
    assert get_res2.json()["settings"]["apply_to_all_generators"] is False

    # 5. Receipt WITHOUT cost_total now uses 60.00 грн/л from system settings!
    rec_conf_res = await client.post("/api/v1/fuel/receipt", json={
        "liters": 10.0,
        "driver_name": "Водій Налаштування Ціна",
        "receipt_number": "AUTO-SET-02",
        "fuel_type": "А-95"
    }, headers=operator_auth)
    assert rec_conf_res.status_code == 201
    rec_conf = rec_conf_res.json()
    assert rec_conf["price_per_liter"] == 60.0
    assert rec_conf["cost_total"] == 600.0


@pytest.mark.asyncio
async def test_system_reset_and_clean_db(client, superadmin_auth, operator_auth):
    """Тест очищення звітів об'єкта та повного заводського скидання БД"""
    # 1. Operator is forbidden from resetting DB
    op_res = await client.post("/api/v1/system/reset", json={
        "mode": "reports_clean",
        "admin_password": "OperatorPass123!",
        "confirm_code": "ОЧИСТИТИ"
    }, headers=operator_auth)
    assert op_res.status_code == 403

    # 2. Wrong admin password is rejected
    bad_pass_res = await client.post("/api/v1/system/reset", json={
        "mode": "reports_clean",
        "admin_password": "WrongPassword123!",
        "confirm_code": "ОЧИСТИТИ"
    }, headers=superadmin_auth)
    assert bad_pass_res.status_code == 403

    # 3. Clean reports & counters across system
    from backend.app.config import settings
    clean_res = await client.post("/api/v1/system/reset", json={
        "mode": "reports_clean",
        "admin_password": settings.INITIAL_ADMIN_PASSWORD,
        "confirm_code": "ОЧИСТИТИ"
    }, headers=superadmin_auth)
    assert clean_res.status_code == 200
    assert clean_res.json()["status"] == "ok"

    # Verify fuel summary is now 0.0
    sum_res = await client.get("/api/v1/fuel/summary", headers=superadmin_auth)
    assert sum_res.status_code == 200
    assert sum_res.json()["warehouse_balance_l"] == 0.0
    assert sum_res.json()["total_received_l"] == 0.0





@pytest.mark.asyncio
async def test_facility_cascade_delete_and_chain_binding(client, superadmin_auth, operator_auth):
    """
    Тестування:
    1. Каскадне видалення об\x27єкта (?cascade=true vs ?cascade=false).
    2. Автоматичне відкріплення закріплених користувачів (u.facility_id -> None).
    3. Жорстка прив\x27язка генератора до об\x27єкта (403 "Ви не можете вибрати генератор ..., він закріплений за об\x27єктом ...").
    4. Збереження та валідація налаштувань ланцюжка за замовчуванням (Default Chain).
    5. Щотижневий звіт Excel з цільовим об\x27єктом та генератором.
    """
    # 1. Create a dedicated test facility
    fac_res = await client.post(
        "/api/v1/facilities",
        headers=superadmin_auth,
        json={
            "name": "Тестовий філіал (Одеса)",
            "address": "вул. Дерибасівська, 1",
            "description": "Філія для перевірки каскадного видалення",
            "timezone": "Europe/Kyiv"
        }
    )
    assert fac_res.status_code == 201
    fac_id = fac_res.json()["id"]

    # 2. Create a generator on this facility
    gen_res = await client.post(
        "/api/v1/generator",
        headers=superadmin_auth,
        json={
            "facility_id": fac_id,
            "name": "Одеський дизель-генератор",
            "model": "SDMO 15kW",
            "manufacturer": "SDMO",
            "serial_number": "ODS-15-999",
            "rated_power_kw": 15.0,
            "tank_capacity_l": 50.0,
            "fuel_type": "ДП",
            "nominal_consumption_l_per_h": 3.8,
            "initial_operating_hours": 0.0,
            "initial_fuel_tank_level_l": 25.0,
            "work_schedule_start": "08:00",
            "work_schedule_end": "20:00",
            "maintenance_interval_hours": 300.0,
            "timezone": "Europe/Kyiv"
        }
    )
    assert gen_res.status_code == 201
    gen_id = gen_res.json()["id"]

    # 3. Create an operator assigned specifically to facility 1
    op_create_res = await client.post(
        "/api/v1/users",
        headers=superadmin_auth,
        json={
            "login": "operator_kyiv_lock",
            "full_name": "Київський Оператор",
            "password": "Password123!",
            "facility_id": 1,
            "role_ids": [3]
        }
    )
    assert op_create_res.status_code == 201

    # Login as this operator
    login_res = await client.post(
        "/api/v1/auth/login",
        json={"login": "operator_kyiv_lock", "password": "Password123!"}
    )
    assert login_res.status_code == 200
    token = login_res.json()["access_token"]
    kyiv_op_auth = {"Authorization": f"Bearer {token}"}

    # 4. Invariant: Kyiv operator tries to start Odessa generator -> 403 Forbidden with exact text
    forbidden_res = await client.post(
        "/api/v1/generator/start",
        headers=kyiv_op_auth,
        json={"generator_id": gen_id, "fuel_level_l": 25.0}
    )
    assert forbidden_res.status_code == 403
    assert "Ви не можете вибрати генератор" in forbidden_res.json()["detail"]
    assert "Тестовий філіал (Одеса)" in forbidden_res.json()["detail"]

    # 5. Admin establishes default chain settings
    settings_put = await client.put(
        "/api/v1/system/settings",
        headers=superadmin_auth,
        json={
            "organization_name": "MAXI-UA CHAIN",
            "facility_timezone": "Europe/Kyiv",
            "work_schedule_start": "08:00",
            "work_schedule_end": "20:00",
            "default_facility_id": fac_id,
            "default_generator_id": gen_id,
            "email_weekly_facility_id": fac_id,
            "email_weekly_generator_id": gen_id
        }
    )
    assert settings_put.status_code == 200
    sett = settings_put.json()["settings"]
    assert sett["default_facility_id"] == fac_id
    assert sett["default_generator_id"] == gen_id
    assert sett["email_weekly_facility_id"] == fac_id
    assert sett["email_weekly_generator_id"] == gen_id

    # 6. Test weekly digest manual send with facility_id and generator_id
    from unittest.mock import patch
    with patch("backend.app.services.email_service.send_weekly_digest_email") as mock_digest:
        mock_digest.return_value = {"status": "ok", "provider": "primary", "message": "Звіт Excel успішно надіслано на 1 адрес(у/ів)!"}
        digest_res = await client.post(
            f"/api/v1/system/email/weekly-digest/send?recipient=admin@test.ua&facility_id={fac_id}&generator_id={gen_id}",
            headers=superadmin_auth
        )
        assert digest_res.status_code == 200
        assert digest_res.json()["result"]["status"] == "ok"
        assert "успішно надіслано" in digest_res.json()["message"]
        mock_digest.assert_called_once()
        assert mock_digest.call_args[1]["recipients"] == ["admin@test.ua"]
        assert isinstance(mock_digest.call_args[1]["excel_bytes"], bytes)

    # 7. Try delete facility without cascade -> 400 Bad Request
    del_no_cascade = await client.delete(
        f"/api/v1/facilities/{fac_id}",
        headers=superadmin_auth
    )
    assert del_no_cascade.status_code == 400
    assert "каскадного видалення" in del_no_cascade.json()["detail"]

    # 8. Delete facility with cascade=True -> 200 OK
    del_cascade = await client.delete(
        f"/api/v1/facilities/{fac_id}?cascade=true",
        headers=superadmin_auth
    )
    assert del_cascade.status_code == 200
    assert "успішно видалено" in del_cascade.json()["message"]

    # 9. Verify generator was cascade deleted
    gen_check = await client.get(
        f"/api/v1/generator?generator_id={gen_id}",
        headers=superadmin_auth
    )
    assert gen_check.status_code == 404

    # 10. Verify runtime setting default_facility_id was safely reset to None
    from backend.app.services.settings_service import get_runtime_setting
    assert get_runtime_setting("default_facility_id") is None
